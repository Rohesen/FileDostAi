"""
Correctness tests for converters.py: each conversion must carry the actual content
across, not just produce a valid file.  Run with:  python3 -m pytest tests
"""

import io
import os
import sys
import unicodedata

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "server"))

import pymupdf  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle  # noqa: E402

import converters  # noqa: E402
from pdf_engine import ToolUnavailableError  # noqa: E402


def all_text(pdf_bytes: bytes) -> str:
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    a = "\n".join(page.get_text() for page in doc)
    b = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(pdf_bytes)).pages)
    # Story output uses ligature glyphs ("fi" -> U+FB01); NFKC turns them back into letters.
    return unicodedata.normalize("NFKC", a + "\n" + b)


def page_count(pdf_bytes: bytes) -> int:
    return pymupdf.open(stream=pdf_bytes, filetype="pdf").page_count


def simple_pdf(pages=2, prefix="Sample page") -> bytes:
    doc = pymupdf.open()
    for i in range(1, pages + 1):
        page = doc.new_page(width=612, height=792)
        page.insert_text((72, 72), f"{prefix} {i}", fontsize=14)
        page.insert_text((72, 100), "The quick brown fox jumps over the lazy dog.", fontsize=11)
    return doc.tobytes()


def table_pdf() -> bytes:
    buf = io.BytesIO()
    data = [["Item", "Qty", "Price"],
            ["Apple", "12", "1,234.50"],
            ["Pear", "7", "3.25"],
            ["Zip code", "00123", "15%"]]
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(buf).build([t])
    return buf.getvalue()


def private_pdf(path) -> bytes:
    """A PDF full of private data, saved with one incremental update."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Contact jane.doe@example.com for details.")
    page.insert_text((72, 96), "Phone: (555) 019-2834")
    page.insert_text((72, 120), "Card 4111 1111 1111 1111")
    page.insert_text((72, 144), "IBAN GB82 WEST 1234 5698 7654 32")
    page.insert_text((72, 300), "invisible secret words", render_mode=3)
    page.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(72, 160, 200, 175),
                      "uri": "https://tracker.example.net/pixel"})
    annot = page.add_text_annot((100, 400), "Please fix this")
    annot.set_info(title="Bob Reviewer")
    annot.update()
    doc.set_metadata({"author": "Alice Author", "creator": "Secret Writer 9", "producer": "Leaky PDF"})
    doc.set_xml_metadata("<x:xmpmeta xmlns:x='adobe:ns:meta/'><rdf:RDF "
                         "xmlns:rdf='http://www.w3.org/1999/02/22-rdf-syntax-ns#'/></x:xmpmeta>")
    doc.embfile_add("payroll.txt", b"salary data", filename="payroll.txt")
    js = doc.get_new_xref()
    doc.update_object(js, "<< /S /JavaScript /JS (app.alert\\(1\\)) >>")
    doc.xref_set_key(doc.pdf_catalog(), "OpenAction", f"{js} 0 R")
    fn = str(path / "private.pdf")
    doc.save(fn)
    doc.close()
    d = pymupdf.open(fn)
    d.set_metadata({"author": "Alice Author", "title": "Second draft"})
    d.saveIncr()
    d.close()
    with open(fn, "rb") as f:
        return f.read()


# ---------- Office ----------

needs_office = pytest.mark.skipif(not converters.office_available(), reason="LibreOffice / MS Office not installed")


@needs_office
def test_docx_to_pdf_keeps_text():
    from docx import Document
    d = Document()
    d.add_heading("Quarterly Report Zebra", 1)
    d.add_paragraph("Revenue grew by 42 percent in the northern region.")
    buf = io.BytesIO()
    d.save(buf)
    pdf = converters.office_to_pdf(buf.getvalue(), "My Report (final).docx")
    text = all_text(pdf)
    assert "Quarterly Report Zebra" in text
    assert "Revenue grew by 42 percent" in text


@needs_office
def test_xlsx_to_pdf_keeps_cells():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.column_dimensions["A"].width = 20
    ws.append(["Fruit", "Quantity"])
    ws.append(["Mangosteen", 17])
    ws.append(["Durian", 4])
    buf = io.BytesIO()
    wb.save(buf)
    text = all_text(converters.office_to_pdf(buf.getvalue(), "stock.xlsx"))
    for word in ("Fruit", "Mangosteen", "Durian", "17"):
        assert word in text


@needs_office
def test_pptx_to_pdf_one_page_per_slide():
    from pptx import Presentation
    prs = Presentation()
    for title in ("First slide Alpha", "Second slide Beta"):
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = title
    buf = io.BytesIO()
    prs.save(buf)
    pdf = converters.office_to_pdf(buf.getvalue(), "deck.pptx")
    assert page_count(pdf) == 2
    text = all_text(pdf)
    assert "First slide Alpha" in text and "Second slide Beta" in text


@needs_office
def test_office_garbage_is_value_error():
    with pytest.raises(ValueError):
        converters.office_to_pdf(b"\x00\x01 definitely not a word file" * 50, "broken.docx")


def test_office_rejects_unknown_extension_and_empty():
    with pytest.raises(ValueError):
        converters.office_to_pdf(b"hello", "virus.exe")
    with pytest.raises(ValueError):
        converters.office_to_pdf(b"", "empty.docx")


def test_office_without_suite(monkeypatch):
    monkeypatch.setattr(converters, "find_soffice", lambda: None)
    monkeypatch.setattr(converters, "_ms_office_apps", lambda: {"Word": False, "Excel": False, "PowerPoint": False})
    assert converters.office_available() is False
    with pytest.raises(ToolUnavailableError):
        converters.office_to_pdf(b"PK\x03\x04 fake", "report.docx")
    # Text and CSV still work without an office suite.
    assert "plain words survive" in all_text(converters.office_to_pdf(b"plain words survive\nline two", "notes.txt"))
    assert "Kumquat" in all_text(converters.office_to_pdf(b"name,qty\nKumquat,3\n", "list.csv"))


# ---------- Markdown / HTML / CSV ----------

def test_markdown_renders_structure():
    md = ("# Annual Plan\n\nSome **bold** and *italic* text with a [link](https://example.com).\n\n"
          "## Budget Table\n\n| Team | Budget |\n|------|--------|\n| Ops | 5000 |\n| R&D | 9000 |\n\n"
          "```python\nprint('hello code')\n```\n\n- first item\n- second item\n")
    pdf = converters.markdown_to_pdf(md)
    text = all_text(pdf)
    for s in ("Annual Plan", "Budget Table", "Ops", "9000", "print('hello code')", "first item", "bold"):
        assert s in text
    assert "**" not in text and "|---" not in text and "```" not in text
    page = pymupdf.open(stream=pdf, filetype="pdf")[0]
    assert page.rect.width == pytest.approx(595, abs=1)  # A4
    assert any(l.get("uri") == "https://example.com" for l in page.get_links())
    fonts = {f[3] for f in page.get_fonts()}
    assert any("Bold" in f for f in fonts)
    assert any("Cour" in f or "Mono" in f for f in fonts)


def test_markdown_long_input_is_multi_page():
    md = "# Long\n\n" + "\n\n".join(f"Paragraph number {i} with some filler text to take space." for i in range(300))
    pdf = converters.markdown_to_pdf(md)
    assert page_count(pdf) > 3
    text = all_text(pdf)
    assert "Paragraph number 0 " in text and "Paragraph number 299 " in text


def test_markdown_fallback_converter():
    html = converters._simple_markdown("# Head\n\n**b** _i_ `c<d`\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n1. one\n2. two\n")
    assert "<h1>Head</h1>" in html and "<b>b</b>" in html and "<i>i</i>" in html
    assert "<code>c&lt;d</code>" in html and "<table>" in html and "<td>2</td>" in html and "<ol>" in html
    pdf = converters._story_to_pdf(html)
    assert "Head" in all_text(pdf)


def test_markdown_empty_is_value_error():
    with pytest.raises(ValueError):
        converters.markdown_to_pdf("   \n ")


def test_html_strips_scripts_and_remote_resources():
    html = ("<html><head><script>document.write('INJECTED SCRIPT TEXT')</script>"
            "<link rel='stylesheet' href='http://evil.example/x.css'></head>"
            "<body><h1>Visible Heading</h1><p onclick='steal()'>Body text here</p>"
            "<img src='http://evil.example/track.png'><iframe src='http://evil.example'>frame text</iframe>"
            "<table><tr><th>K</th><th>V</th></tr><tr><td>alpha</td><td>beta</td></tr></table></body></html>")
    clean = converters._sanitize_html(html)
    assert "evil.example" not in clean and "onclick" not in clean and "<script" not in clean.lower()
    text = all_text(converters.html_to_pdf(html))
    assert "Visible Heading" in text and "Body text here" in text and "alpha" in text
    assert "INJECTED" not in text and "frame text" not in text


def test_csv_to_pdf_cells_and_layout():
    rows = ["Name;City;Score"] + [f"Person {i};Town <{i}>;{i * 3}" for i in range(400)]
    pdf = converters.csv_to_pdf("\n".join(rows).encode("utf-8"))
    text = all_text(pdf)
    assert "Person 0" in text and "Person 399" in text and "Town <17>" in text and "1197" in text
    assert "Name;City" not in text  # semicolon delimiter was detected
    assert page_count(pdf) > 1
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    assert doc[0].rect.width < doc[0].rect.height  # 3 columns -> portrait
    wide = converters.csv_to_pdf(("\t".join(f"col{i}" for i in range(10)) + "\n" + "\t".join("x" * 10)).encode())
    wdoc = pymupdf.open(stream=wide, filetype="pdf")
    assert wdoc[0].rect.width > wdoc[0].rect.height  # many columns -> landscape
    assert "col9" in all_text(wide)


def test_csv_latin1_and_empty():
    assert "Café" in all_text(converters.csv_to_pdf("drink,price\nCafé,2\n".encode("cp1252")))
    with pytest.raises(ValueError):
        converters.csv_to_pdf(b"  \n")


# ---------- PDF -> Excel ----------

def test_pdf_to_excel_recovers_table_with_numbers():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(converters.pdf_to_excel(table_pdf())))
    assert wb.sheetnames == ["Page 1"]
    ws = wb["Page 1"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ["Item", "Qty", "Price"]
    assert rows[1] == ["Apple", 12, 1234.5]
    assert isinstance(rows[1][1], int) and isinstance(rows[1][2], float)
    assert rows[2] == ["Pear", 7, 3.25]
    assert rows[3][1] == "00123"  # leading zeros stay text
    assert rows[3][2] == pytest.approx(0.15)
    assert ws["A1"].font.b
    assert ws.column_dimensions["A"].width >= 8


def test_pdf_to_excel_without_tables_gives_text():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(converters.pdf_to_excel(simple_pdf(2))))
    ws = wb.active
    values = [r for r in ws.iter_rows(values_only=True)]
    assert values[0] == ("Page", "Text")
    assert (1, "Sample page 1") in values and (2, "Sample page 2") in values


def test_number_parsing():
    f = converters._to_number
    assert f("-1,200") == -1200 and f("(5)") == -5 and f("$3.50") == 3.5 and f("12%") == pytest.approx(0.12)
    assert f("1.2.3") == "1.2.3" and f("007") == "007" and f("12345678901234567") == "12345678901234567"


def test_bad_pdf_inputs_are_value_errors():
    for fn in (converters.pdf_to_excel, converters.pdf_to_pptx, converters.privacy_scan):
        with pytest.raises(ValueError):
            fn(b"not a pdf at all")
        with pytest.raises(ValueError):
            fn(b"")
    locked = pymupdf.open(stream=simple_pdf(1), filetype="pdf").tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    with pytest.raises(ValueError, match="password"):
        converters.pdf_to_excel(locked)


# ---------- PDF -> PowerPoint ----------

def test_pdf_to_pptx_slides_and_notes():
    from pptx import Presentation
    pdf = simple_pdf(3, prefix="Notes marker")
    prs = Presentation(io.BytesIO(converters.pdf_to_pptx(pdf)))
    assert len(prs.slides) == 3
    assert prs.slide_height / prs.slide_width == pytest.approx(792 / 612, rel=0.01)
    for i, slide in enumerate(prs.slides, start=1):
        pics = [s for s in slide.shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
        assert len(pics) == 1
        assert pics[0].width == pytest.approx(prs.slide_width, abs=2)
        assert f"Notes marker {i}" in slide.notes_slide.notes_text_frame.text


def test_pdf_to_pptx_landscape_page():
    from pptx import Presentation
    doc = pymupdf.open()
    doc.new_page(width=842, height=595).insert_text((72, 72), "wide")
    prs = Presentation(io.BytesIO(converters.pdf_to_pptx(doc.tobytes())))
    assert prs.slide_width > prs.slide_height


# ---------- PDF -> Word ----------

def test_pdf_to_docx_layout():
    pytest.importorskip("pdf2docx")
    from docx import Document
    d = Document(io.BytesIO(converters.pdf_to_docx(table_pdf())))
    cells = [c.text.strip() for t in d.tables for row in t.rows for c in row.cells]
    assert "Apple" in cells and "1,234.50" in cells
    d2 = Document(io.BytesIO(converters.pdf_to_docx(simple_pdf(2))))
    text = "\n".join(p.text for p in d2.paragraphs)
    assert "Sample page 1" in text and "Sample page 2" in text


def test_pdf_to_docx_missing_library(monkeypatch):
    import builtins
    real_import = builtins.__import__

    def fake_import(name, *a, **k):
        if name == "pdf2docx" or name.startswith("pdf2docx."):
            raise ImportError("no pdf2docx")
        return real_import(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(ToolUnavailableError):
        converters.pdf_to_docx(simple_pdf(1))


# ---------- Repair ----------

def test_repair_broken_xref():
    good = simple_pdf(3)
    bad = good.replace(b"startxref", b"startxxxx").replace(b"xref", b"xxxx")
    result = converters.repair_pdf(bad)
    assert page_count(result["bytes"]) == 3
    assert "Sample page 3" in all_text(result["bytes"])
    assert result["report"] and all(isinstance(s, str) for s in result["report"])
    # The repaired file opens cleanly, without needing recovery.
    import pikepdf
    with pikepdf.open(io.BytesIO(result["bytes"])) as p:
        assert len(p.pages) == 3 and not p.get_warnings()


def test_repair_truncated_tail_and_junk():
    good = pymupdf.open(stream=simple_pdf(3), filetype="pdf").tobytes(deflate=False, use_objstms=False)
    truncated = good[: int(len(good) * 0.85)]
    result = converters.repair_pdf(truncated)
    assert page_count(result["bytes"]) >= 1
    junky = b"GARBAGE-HEADER" + good + b"\x00trailing junk" * 20
    result2 = converters.repair_pdf(junky)
    assert page_count(result2["bytes"]) == 3
    assert any("junk" in s for s in result2["report"])


def test_repair_healthy_and_hopeless():
    result = converters.repair_pdf(simple_pdf(2))
    assert page_count(result["bytes"]) == 2
    assert any("No damage" in s for s in result["report"])
    with pytest.raises(ValueError):
        converters.repair_pdf(b"this is just text, not a pdf")
    with pytest.raises(ValueError):
        converters.repair_pdf(b"%PDF-1.7\n" + b"\x00" * 200)


# ---------- Privacy ----------

def test_privacy_scan_finds_everything(tmp_path):
    report = converters.privacy_scan(private_pdf(tmp_path))
    kinds = {f["kind"]: f for f in report["findings"]}
    for k in ("metadata-author", "xmp", "javascript", "attachments", "annotations", "email", "phone",
              "card", "iban", "history", "links", "hidden-text"):
        assert k in kinds, k
    assert "Alice Author" in kinds["metadata-author"]["detail"]
    assert "payroll.txt" in kinds["attachments"]["detail"]
    assert "Bob Reviewer" in kinds["annotations"]["detail"]
    assert "invisible secret" in kinds["hidden-text"]["detail"]
    assert "tracker.example.net" in kinds["links"]["detail"]
    assert kinds["javascript"]["severity"] == "high" and kinds["card"]["severity"] == "high"
    assert kinds["email"]["fix"] == "redact" and kinds["metadata-author"]["fix"] == "remove-metadata"
    for f in report["findings"]:
        assert f["severity"] in ("high", "medium", "low")
        assert f["fix"] in ("remove-metadata", "redact", "flatten", None)
        assert f["title"] and f["detail"]
    assert "jane.doe@example.com" in report["found_terms"]
    assert "(555) 019-2834" in report["found_terms"]
    assert "4111 1111 1111 1111" in report["found_terms"]
    assert 0 <= report["score"] < 50


def test_privacy_scan_clean_file_scores_high():
    report = converters.privacy_scan(simple_pdf(1))
    assert report["score"] >= 90
    assert not [f for f in report["findings"] if f["severity"] != "low"]
    assert report["found_terms"] == []


def test_privacy_scan_ignores_non_luhn_numbers():
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Order 1234 5678 9012 3456 shipped")
    report = converters.privacy_scan(doc.tobytes())
    assert "card" not in {f["kind"] for f in report["findings"]}


def test_strip_private_data_removes_it(tmp_path):
    src = private_pdf(tmp_path)
    before = converters.privacy_scan(src)
    out = converters.strip_private_data(src)
    after = converters.privacy_scan(out)
    kinds = {f["kind"] for f in after["findings"]}
    for k in ("metadata-author", "metadata", "xmp", "javascript", "attachments", "history"):
        assert k not in kinds, k
    assert after["score"] > before["score"]
    doc = pymupdf.open(stream=out, filetype="pdf")
    assert not doc.metadata.get("author") and not doc.metadata.get("producer")
    assert doc.embfile_count() == 0
    assert not (doc.get_xml_metadata() or "").strip()
    assert b"app.alert" not in out and b"salary data" not in out and b"Alice Author" not in out
    assert out.count(b"%%EOF") == 1
    authors = [a.info.get("title") for a in doc[0].annots()]
    assert "Bob Reviewer" not in authors
    # Page content and links are untouched.
    assert "jane.doe@example.com" in doc[0].get_text()
    assert any(l.get("uri") for l in doc[0].get_links())


def test_strip_keeps_form_field_names():
    doc = pymupdf.open()
    page = doc.new_page()
    w = pymupdf.Widget()
    w.field_name = "full_name"
    w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    w.rect = pymupdf.Rect(72, 100, 372, 130)
    w.field_value = "Jane Q Public"
    page.add_widget(w)
    out = converters.strip_private_data(doc.tobytes())
    names = [x.field_name for x in pymupdf.open(stream=out, filetype="pdf")[0].widgets()]
    assert names == ["full_name"]


# ---------- PDF/A ----------

needs_gs = pytest.mark.skipif(not converters.find_ghostscript(), reason="Ghostscript not installed")


@needs_gs
def test_pdf_to_pdfa_identification():
    out = converters.pdf_to_pdfa(simple_pdf(2))
    doc = pymupdf.open(stream=out, filetype="pdf")
    assert doc.page_count == 2
    xmp = doc.get_xml_metadata()
    assert "pdfaid" in xmp and "2" in xmp.split("pdfaid:part")[1][:5]
    assert "Sample page 2" in all_text(out)
    import pikepdf
    with pikepdf.open(io.BytesIO(out)) as p:
        assert "/OutputIntents" in p.Root


def test_pdfa_without_ghostscript(monkeypatch):
    monkeypatch.setattr(converters, "find_ghostscript", lambda: None)
    with pytest.raises(ToolUnavailableError):
        converters.pdf_to_pdfa(simple_pdf(1))


# ----- wide spreadsheets keep every column on the page -----

def _wide_rows(n_cols=18, n_rows=40):
    import random
    random.seed(3)
    head = ["experiment", "split"] + [f"col{i}" for i in range(n_cols - 3)] + ["LAST_COLUMN"]
    return head, [[f"exp{r}", "train"] + [random.random() for _ in range(n_cols - 2)] for r in range(n_rows)]


def test_wide_csv_fits_every_column_on_first_page():
    import csv as csv_mod
    head, rows = _wide_rows()
    buf = io.StringIO()
    w = csv_mod.writer(buf)
    w.writerow(head)
    for r in rows:
        w.writerow([repr(v) if isinstance(v, float) else v for v in r])
    pdf = converters.csv_to_pdf(buf.getvalue().encode())
    doc = pymupdf.open(stream=pdf)
    first = doc[0].get_text()
    assert "experiment" in first and "LAST_COLUMN" in first
    # header repeats on later pages, and every word sits inside the page
    for page in doc:
        assert "LAST_COLUMN" in page.get_text()
        for x0, y0, x1, y1, *_ in page.get_text("words"):
            assert 0 <= x0 and x1 <= page.rect.width + 0.5


def test_long_floats_are_rounded_like_excel():
    pdf = converters.csv_to_pdf(b"a,b\n0.12345678901234567,2\n")
    text = pymupdf.open(stream=pdf)[0].get_text()
    assert "0.123456789" in text and "0.12345678901234567" not in text


def test_xlsx_page_setup_fits_columns_and_keeps_data():
    import openpyxl
    import zipfile
    head, rows = _wide_rows()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(head)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    fixed = converters.fit_xlsx_columns(out.getvalue())
    xml = zipfile.ZipFile(io.BytesIO(fixed)).read("xl/worksheets/sheet1.xml").decode()
    assert 'fitToPage="1"' in xml and 'fitToWidth="1"' in xml and 'fitToHeight="0"' in xml and 'orientation="landscape"' in xml
    again = openpyxl.load_workbook(io.BytesIO(fixed))
    assert again.active.max_column == 18 and again.active["R1"].value == "LAST_COLUMN"


def test_fit_sheet_xml_handles_existing_page_setup():
    xml = ('<worksheet xmlns="x"><sheetPr><tabColor rgb="FF0000"/></sheetPr><dimension ref="A1:C3"/><sheetData/>'
           '<pageMargins left="1" right="1" top="1" bottom="1" header="0" footer="0"/>'
           '<pageSetup scale="80" fitToWidth="3" orientation="portrait"/><drawing r:id="rId1"/></worksheet>')
    out = converters._fit_sheet_xml(xml)
    assert '<tabColor rgb="FF0000"/><pageSetUpPr fitToPage="1"/></sheetPr>' in out
    assert 'scale=' not in out and out.count("fitToWidth") == 1 and 'orientation="portrait"' in out
    assert out.index("<pageSetup") < out.index("<drawing")


def test_xlsx_without_office_suite_draws_all_columns(monkeypatch):
    import openpyxl
    head, rows = _wide_rows()
    wb = openpyxl.Workbook()
    wb.active.append(head)
    for r in rows:
        wb.active.append(r)
    out = io.BytesIO()
    wb.save(out)
    monkeypatch.setattr(converters, "find_soffice", lambda: None)
    monkeypatch.setattr(converters, "_ms_office_apps", lambda: {})
    pdf = converters.office_to_pdf(out.getvalue(), "results.xlsx")
    assert "LAST_COLUMN" in pymupdf.open(stream=pdf)[0].get_text()


@pytest.mark.skipif(not converters.find_soffice(), reason="LibreOffice not installed")
def test_wide_xlsx_through_libreoffice_fits_page_width():
    import openpyxl
    head, rows = _wide_rows()
    wb = openpyxl.Workbook()
    wb.active.append(head)
    for r in rows:
        wb.active.append(r)
    out = io.BytesIO()
    wb.save(out)
    pdf = converters.office_to_pdf(out.getvalue(), "results.xlsx")
    assert "LAST_COLUMN" in pymupdf.open(stream=pdf)[0].get_text()
