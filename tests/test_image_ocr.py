"""Image to Text and JPG to Word (image_ocr.py and their endpoints)."""

import io
import os
import sys

import docx
import pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import image_ocr  # noqa: E402
from main import app  # noqa: E402
from pdf_engine import PDFEngine  # noqa: E402

needs_ocr = pytest.mark.skipif(not PDFEngine.ocr_available(), reason="tesseract not installed")
client = TestClient(app)


def _font(size, bold=False):
    for path in ("/usr/share/fonts/truetype/dejavu/DejaVuSans%s.ttf" % ("-Bold" if bold else ""),
                 "C:/Windows/Fonts/arial%s.ttf" % ("bd" if bold else "")):
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def page_photo(fmt="JPEG", rotate_exif=False) -> bytes:
    img = Image.new("RGB", (1200, 700), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 40), "Exam Admit Card", font=_font(56, True), fill="black")
    y = 150
    for line in ["The quick brown fox jumps over the lazy dog and keeps",
                 "running through the green fields until the evening sun",
                 "finally sets behind the distant hills of the valley."]:
        d.text((60, y), line, font=_font(34), fill="black")
        y += 48
    y += 60
    for line in ["Name: Rohan Kumar", "Roll No: 123456"]:
        d.text((60, y), line, font=_font(34), fill="black")
        y += 48
    out = io.BytesIO()
    if rotate_exif:
        # Stored sideways with an EXIF "rotate 90" tag, like many phone photos.
        img = img.transpose(Image.Transpose.ROTATE_90)
        exif = Image.Exif()
        exif[0x0112] = 6
        img.save(out, fmt, exif=exif)
    else:
        img.save(out, fmt)
    return out.getvalue()


@needs_ocr
def test_image_to_text_reads_the_words():
    r = image_ocr.images_to_text([("card.jpg", page_photo())])["results"][0]
    assert "quick brown fox" in r["text"]
    assert "Roll No: 123456" in r["text"]
    assert r["words"] > 30


@needs_ocr
def test_sideways_phone_photo_is_turned_upright_first():
    r = image_ocr.images_to_text([("card.jpg", page_photo(rotate_exif=True))])["results"][0]
    assert "brown fox jumps" in r["text"] and "Exam Admit Card" in r["text"]


@needs_ocr
def test_jpg_to_word_keeps_paragraphs_and_the_heading():
    data = image_ocr.images_to_docx([("card.png", page_photo("PNG")), ("card2.jpg", page_photo())])
    paras = [p for p in docx.Document(io.BytesIO(data)).paragraphs if p.text.strip()]
    texts = [p.text for p in paras]
    assert texts[0] == "Exam Admit Card" and paras[0].runs[0].bold
    # The three wrapped lines become one paragraph.
    assert any(t.startswith("The quick brown fox") and t.endswith("valley.") for t in texts)
    assert sum(t == "Exam Admit Card" for t in texts) == 2  # both pictures


@needs_ocr
def test_endpoints():
    res = client.post("/api/image-to-text", files=[("files", ("a.jpg", page_photo(), "image/jpeg"))], data={"language": "eng"})
    assert res.status_code == 200 and "quick brown fox" in res.json()["results"][0]["text"]
    res = client.post("/api/image-to-word", files=[("files", ("my card.jpg", page_photo(), "image/jpeg"))])
    assert res.status_code == 200
    assert "my card.docx" in res.headers["content-disposition"]
    assert "quick brown fox" in "\n".join(p.text for p in docx.Document(io.BytesIO(res.content)).paragraphs)


@needs_ocr
def test_bad_input_gets_a_clear_message():
    res = client.post("/api/image-to-text", files=[("files", ("a.jpg", b"not a picture", "image/jpeg"))])
    assert res.status_code == 400 and "isn't a picture" in res.json()["detail"]
    res = client.post("/api/image-to-text", files=[("files", ("a.jpg", page_photo(), "image/jpeg"))], data={"language": "eng; rm"})
    assert res.status_code == 400
    blank = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(blank, "PNG")
    res = client.post("/api/image-to-word", files=[("files", ("blank.png", blank.getvalue(), "image/png"))])
    assert res.status_code == 400 and "No text was found" in res.json()["detail"]
    files = [("files", (f"{i}.png", blank.getvalue(), "image/png")) for i in range(image_ocr.MAX_IMAGES + 1)]
    assert client.post("/api/image-to-text", files=files).status_code == 400
