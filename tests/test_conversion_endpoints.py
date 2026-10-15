"""The conversion, repair and privacy endpoints, end to end through FastAPI."""

import io
import zipfile

import pymupdf
import pytest
from fastapi.testclient import TestClient

import converters
import main

client = TestClient(main.app)


def text_pdf(text="Hello from UltraPDF") -> bytes:
    d = pymupdf.open()
    d.new_page().insert_text((72, 100), text, fontsize=14)
    d.set_metadata({"author": "Jane Secret"})
    return d.tobytes()


def post(path, name, data, **form):
    return client.post(path, files={"file": (name, data)}, data=form)


def pdf_text(data):
    import unicodedata
    return unicodedata.normalize("NFKC", "".join(p.get_text() for p in pymupdf.open(stream=data, filetype="pdf")))


def test_markdown_to_pdf_with_unicode_filename():
    r = post("/api/to-pdf", "Notizen über Pläne.md", b"# Plan\n\n- first item\n- second item\n")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert "Notizen%20%C3%BCber%20Pl%C3%A4ne.pdf" in r.headers["content-disposition"]
    assert "first item" in pdf_text(r.content)


def test_csv_and_html_and_text_to_pdf():
    assert "Widget" in pdf_text(post("/api/to-pdf", "t.csv", b"name,price\nWidget,3.50\n").content)
    r = post("/api/to-pdf", "p.html", b"<h1>Title</h1><script>alert(1)</script><p>Body text</p>")
    assert "Body text" in pdf_text(r.content) and "alert" not in pdf_text(r.content)
    assert "plain words" in pdf_text(post("/api/to-pdf", "a.txt", b"plain words").content)


@pytest.mark.skipif(not converters.office_available(), reason="no LibreOffice")
def test_word_to_pdf():
    import docx
    buf = io.BytesIO()
    d = docx.Document()
    d.add_heading("Quarterly report", 1)
    d.add_paragraph("Revenue grew nicely.")
    d.save(buf)
    r = post("/api/to-pdf", "report.docx", buf.getvalue())
    assert r.status_code == 200, r.text
    assert "Revenue grew nicely." in pdf_text(r.content)


def test_unknown_type_and_empty_file_are_refused():
    assert post("/api/to-pdf", "x.exe", b"MZ").status_code == 400
    assert post("/api/to-pdf", "x.md", b"").status_code == 400


def test_pdf_to_excel_and_powerpoint():
    r = post("/api/pdf-to-excel", "t.pdf", text_pdf())
    assert r.status_code == 200 and r.headers["content-disposition"].startswith('attachment; filename="t.xlsx"')
    assert zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    r = post("/api/pdf-to-powerpoint", "t.pdf", text_pdf())
    assert r.status_code == 200
    from pptx import Presentation
    deck = Presentation(io.BytesIO(r.content))
    assert len(deck.slides) == 1 and "Hello" in deck.slides[0].notes_slide.notes_text_frame.text


def test_pdf_to_word_returns_docx():
    r = post("/api/pdf-to-docx", "t.pdf", text_pdf())
    assert r.status_code == 200
    import docx
    assert "Hello from UltraPDF" in "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)


def test_repair_reports_and_fixes_a_cut_off_file():
    good = text_pdf()
    broken = good[: int(len(good) * 0.97)]
    r = post("/api/repair", "b.pdf", broken)
    assert r.status_code == 200, r.text
    assert "Hello from UltraPDF" in pdf_text(r.content)
    assert "x-repair-report" in r.headers


def test_privacy_scan_then_strip():
    d = pymupdf.open(stream=text_pdf("Mail jane.doe@example.com today"), filetype="pdf")
    data = d.tobytes()
    scan = post("/api/privacy-scan", "p.pdf", data).json()
    assert scan["score"] < 100 and "jane.doe@example.com" in scan["found_terms"]
    assert any(f["fix"] == "remove-metadata" for f in scan["findings"])
    clean = post("/api/strip-private", "p.pdf", data)
    assert clean.status_code == 200
    assert not pymupdf.open(stream=clean.content, filetype="pdf").metadata.get("author")


def test_bad_pdf_gives_400_not_500():
    for path in ("/api/pdf-to-excel", "/api/pdf-to-powerpoint", "/api/privacy-scan", "/api/strip-private"):
        assert post(path, "x.pdf", b"not a pdf").status_code == 400, path


@pytest.mark.skipif(converters.find_ghostscript() is None, reason="no Ghostscript")
def test_pdfa():
    r = post("/api/pdf-to-pdfa", "t.pdf", text_pdf())
    assert r.status_code == 200 and b"pdfaid" in r.content
