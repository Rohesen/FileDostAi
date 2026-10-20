"""
Reads the text in photos and scans (JPG, PNG, WebP, ...) with Tesseract, through PyMuPDF.
Used by Image to Text (copyable text) and JPG to Word (an editable .docx).
"""

import io
import os
import re
import statistics
from typing import Any, Dict, List, Tuple

import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from PIL import Image, ImageOps

from pdf_engine import ToolUnavailableError, _find_tessdata

MAX_IMAGES = 20
MAX_OCR_PIXELS = 40_000_000  # what Tesseract gets to read, after any enlarging


def _check_language(language: str) -> str:
    if not re.fullmatch(r"[a-z_]{3,8}(\+[a-z_]{3,8})*", language or ""):
        raise ValueError("Language must be Tesseract codes like 'eng' or 'eng+fra'.")
    tessdata = _find_tessdata()
    if tessdata is None:
        raise ToolUnavailableError("OCR engine is not installed on this server (needs Tesseract).")
    missing = [l for l in language.split("+") if not os.path.exists(os.path.join(tessdata, f"{l}.traineddata"))]
    if missing:
        raise ValueError(f"The language pack for '{missing[0]}' isn't installed on this server.")
    return tessdata


def _prepare(name: str, data: bytes) -> Tuple[bytes, int, int, float]:
    """Upright, flattened PNG of the picture, plus the zoom Tesseract should read it at."""
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Image.DecompressionBombError:
        raise
    except Exception:
        raise ValueError(f"{name or 'This file'} isn't a picture this tool can read. Use JPG, PNG, WebP, BMP, GIF or TIFF.")
    img = ImageOps.exif_transpose(img)
    if img.mode in ("RGBA", "LA", "P", "PA"):
        img = img.convert("RGBA")
        flat = Image.new("RGB", img.size, "white")
        flat.paste(img, mask=img.getchannel("A"))
        img = flat
    elif img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    w, h = img.size
    # Tesseract reads best when letters are 25-40 px tall, so small photos and screenshots are enlarged.
    zoom = min(3.0, 2400 / max(w, h)) if max(w, h) < 1800 else 1.0
    zoom = min(zoom, (MAX_OCR_PIXELS / (w * h)) ** 0.5)
    out = io.BytesIO()
    img.save(out, format="PNG")
    return out.getvalue(), w, h, max(zoom, 0.25)


def _read(name: str, data: bytes, language: str, tessdata: str) -> List[List[Dict[str, Any]]]:
    """Paragraphs in reading order; each is a list of lines {text, width, height}."""
    png, w, h, zoom = _prepare(name, data)
    doc = pymupdf.open()
    try:
        page = doc.new_page(width=w, height=h)  # 1 pixel = 1 point, so dpi 72 reads it at full size
        page.insert_image(page.rect, stream=png)
        tp = page.get_textpage_ocr(language=language, dpi=int(round(72 * zoom)), full=True, tessdata=tessdata)
        info = page.get_text("dict", textpage=tp, sort=True)
    finally:
        doc.close()
    lines = []
    for block in info.get("blocks", []):
        for line in block.get("lines", []):
            text = " ".join(s["text"].strip() for s in line.get("spans", []) if s["text"].strip())
            text = re.sub(r"\s+", " ", text).strip()
            if text:
                x0, y0, x1, y1 = line["bbox"]
                lines.append({"text": text, "x0": x0, "y0": y0, "x1": x1, "y1": y1, "width": x1 - x0, "height": y1 - y0})
    return _paragraphs(lines)


def _paragraphs(lines):
    """Tesseract gives one block per line here; lines that sit close together, start at about the
    same place and have the same letter size are put back into paragraphs."""
    paragraphs = []
    for l in lines:
        p = paragraphs[-1] if paragraphs else None
        if p:
            prev = p[-1]
            h = max(prev["height"], l["height"])
            gap = l["y0"] - prev["y1"]
            same_size = 0.75 < l["height"] / max(prev["height"], 1e-6) < 1.33
            if -0.5 * h < gap < 0.6 * h and same_size and abs(l["x0"] - prev["x0"]) < 1.5 * h:
                p.append(l)
                continue
        paragraphs.append([l])
    return paragraphs


def _plain(paragraphs) -> str:
    return "\n\n".join("\n".join(l["text"] for l in p) for p in paragraphs)


def images_to_text(images: List[Tuple[str, bytes]], language: str = "eng") -> Dict[str, Any]:
    if not images:
        raise ValueError("Add at least one picture.")
    if len(images) > MAX_IMAGES:
        raise ValueError(f"Up to {MAX_IMAGES} pictures at a time, please.")
    tessdata = _check_language(language)
    results = []
    for name, data in images:
        text = _plain(_read(name, data, language, tessdata))
        results.append({"name": name, "text": text, "words": len(text.split())})
    return {"results": results, "language": language}


def _join(lines) -> str:
    """Lines of one wrapped paragraph become one line; a hyphen split across lines is mended."""
    out = ""
    for l in lines:
        t = l["text"]
        if out.endswith("-") and t[:1].islower():
            out = out[:-1] + t
        else:
            out = f"{out} {t}" if out else t
    return out


def images_to_docx(images: List[Tuple[str, bytes]], language: str = "eng") -> bytes:
    if not images:
        raise ValueError("Add at least one picture.")
    if len(images) > MAX_IMAGES:
        raise ValueError(f"Up to {MAX_IMAGES} pictures at a time, please.")
    tessdata = _check_language(language)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    found = 0
    for i, (name, data) in enumerate(images):
        paragraphs = _read(name, data, language, tessdata)
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        heights = [l["height"] for p in paragraphs for l in p]
        body = statistics.median(heights) if heights else 0
        for p in paragraphs:
            found += 1
            widest = max(l["width"] for l in p)
            # Wrapped prose (most lines nearly as wide as the widest) becomes one flowing paragraph;
            # lists, addresses and tables keep their line breaks.
            flowing = len(p) > 1 and sum(l["width"] > widest * 0.7 for l in p[:-1]) >= 0.8 * (len(p) - 1)
            para = doc.add_paragraph()
            big = body and statistics.median(l["height"] for l in p) > body * 1.25
            if flowing:
                run = para.add_run(_join(p))
                run.bold = bool(big)
            else:
                for j, l in enumerate(p):
                    run = para.add_run(l["text"])
                    run.bold = bool(big)
                    if j < len(p) - 1:
                        run.add_break()
    if not found:
        raise ValueError("No text was found in this picture. Check the language, and use a sharper, well-lit photo.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
