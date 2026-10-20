"""
UltraPDF converters - server-side conversions that need more than the core engine:
Office documents, Markdown/HTML/CSV to PDF, PDF to Excel/PowerPoint/Word,
PDF repair, privacy scanning/cleaning and PDF/A.

Every function takes bytes (or text) and returns bytes or a dict.  Bad input raises
ValueError with a plain-English message; a missing helper program (LibreOffice,
Ghostscript, pdf2docx) raises ToolUnavailableError so the caller can explain or fall back.
"""

import contextlib
import csv
import glob
import html as html_lib
import io
import logging
import os
import re
import shutil
import signal
import stat
import queue
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pymupdf

from pdf_engine import ToolUnavailableError

# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

A4 = pymupdf.paper_rect("a4")
MARGIN = 50
MAX_PAGES_OUT = 5000

# Characters that are not allowed in XML (docx/xlsx/pptx) files.
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _xml_safe(text: str) -> str:
    return _XML_ILLEGAL.sub("", text or "")


def _open_pdf(pdf: bytes) -> "pymupdf.Document":
    """Opens PDF bytes with PyMuPDF, turning every failure into a friendly ValueError."""
    if not pdf:
        raise ValueError("The file is empty.")
    try:
        doc = pymupdf.open(stream=pdf, filetype="pdf")
    except Exception:
        raise ValueError("This file is not a readable PDF. Try the Repair tool first.")
    if doc.needs_pass:
        doc.close()
        raise ValueError("This PDF is password-protected. Unlock it first.")
    if doc.page_count == 0:
        doc.close()
        raise ValueError("This PDF has no pages.")
    return doc


def _windows_program_dirs() -> List[str]:
    return [d for d in (os.environ.get("ProgramFiles", r"C:\Program Files"),
                        os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)"),
                        os.environ.get("ProgramW6432", ""),
                        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs")) if d]


def _run(cmd: List[str], timeout: int, cwd: Optional[str] = None) -> subprocess.CompletedProcess:
    """Runs a helper program, killing it (and any children) if it takes too long."""
    kwargs: Dict[str, Any] = {}
    if os.name == "nt":
        kwargs["creationflags"] = getattr(subprocess, "CREATE_NO_WINDOW", 0) | getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
    else:
        kwargs["start_new_session"] = True
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, stdin=subprocess.DEVNULL, cwd=cwd, **kwargs)
    try:
        out, err = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        try:
            if os.name == "nt":
                subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)], capture_output=True)
            else:
                os.killpg(proc.pid, signal.SIGKILL)
        except Exception:
            proc.kill()
        proc.communicate()
        raise ValueError("The conversion took too long and was stopped. Try a smaller file.")
    return subprocess.CompletedProcess(cmd, proc.returncode, out, err)


# ---------------------------------------------------------------------------
# 1. Office documents -> PDF
# ---------------------------------------------------------------------------

WORD_EXTS = {".doc", ".docx", ".odt", ".rtf", ".txt"}
EXCEL_EXTS = {".xls", ".xlsx", ".ods", ".csv"}
POWERPOINT_EXTS = {".ppt", ".pptx", ".odp"}
OFFICE_EXTS = WORD_EXTS | EXCEL_EXTS | POWERPOINT_EXTS


def find_soffice() -> Optional[str]:
    """Path to LibreOffice's soffice program, or None."""
    for name in ("soffice", "libreoffice", "soffice.com", "soffice.exe"):
        exe = shutil.which(name)
        if exe:
            return exe
    candidates = []
    if os.name == "nt":
        for root in _windows_program_dirs():
            candidates.append(os.path.join(root, "LibreOffice", "program", "soffice.exe"))
            candidates += glob.glob(os.path.join(root, "LibreOffice*", "program", "soffice.exe"))
    else:
        candidates += ["/Applications/LibreOffice.app/Contents/MacOS/soffice",
                       "/usr/lib/libreoffice/program/soffice", "/opt/libreoffice/program/soffice"]
        candidates += sorted(glob.glob("/opt/libreoffice*/program/soffice"), reverse=True)
    for c in candidates:
        if os.path.isfile(c):
            return c
    return None


def _ms_office_apps() -> Dict[str, bool]:
    """Which Microsoft Office apps can be driven through COM (Windows + pywin32 only)."""
    apps = {"Word": False, "Excel": False, "PowerPoint": False}
    if os.name != "nt":
        return apps
    try:
        import win32com.client  # noqa: F401
        import winreg
    except Exception:
        return apps
    for app in apps:
        try:
            winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, f"{app}.Application\\CLSID"))
            apps[app] = True
        except OSError:
            pass
    return apps


def office_available() -> bool:
    """True when Office files can be converted (LibreOffice, or Microsoft Office on Windows)."""
    return bool(find_soffice()) or any(_ms_office_apps().values())


def _safe_name(filename: str, ext: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(filename or "document").stem)[:60].strip("_") or "document"
    return stem + ext


def private_work_dir() -> str:
    """A folder only this user can open, for LibreOffice profiles and lock files. A fixed name
    in the shared temp folder could be planted in advance by another account on the machine."""
    base = os.path.join(tempfile.gettempdir(), f"ultrapdf-{os.getuid()}" if hasattr(os, "getuid") else "ultrapdf")
    try:
        os.mkdir(base, 0o700)
    except FileExistsError:
        pass
    if hasattr(os, "getuid"):
        st = os.lstat(base)
        if not stat.S_ISDIR(st.st_mode) or st.st_uid != os.getuid() or st.st_mode & 0o077:
            raise RuntimeError(f"{base} is not a private folder owned by this user; refusing to use it.")
    return base


# Settings written into every LibreOffice profile before it runs, so an uploaded document
# can't run macros, recalculate formulas such as WEBSERVICE(), or make the server fetch
# linked images and files from the network or the disk (tested: without these a .docx or
# .odt with a linked picture makes the server request any URL it names).
_LO_LOCKDOWN = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="SecureURL" oor:op="fuse"><value/></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>1</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>1</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="UseOpenCL" oor:op="fuse"><value>false</value></prop></item>
</oor:items>
"""


# One reusable LibreOffice profile per slot: creating a profile is the slowest part of a
# cold start (many seconds on Windows), so conversions reuse a warm one. Jobs may run in
# separate processes, so on Linux/macOS a slot is claimed with a file lock, not just a queue.
_LO_SLOTS = 2
_lo_pool: "queue.Queue[str]" = queue.Queue()
for _i in range(_LO_SLOTS):
    _lo_pool.put(str(_i))


@contextlib.contextmanager
def _lo_profile(wait: float = 180):
    base = private_work_dir()
    try:
        import fcntl
    except ImportError:
        fcntl = None
    if fcntl is None:
        slot = _lo_pool.get()
        try:
            yield os.path.join(base, f"lo_profile_{slot}")
        finally:
            _lo_pool.put(slot)
        return
    deadline = time.monotonic() + wait
    while True:
        for slot in range(_LO_SLOTS):
            fd = os.open(os.path.join(base, f"lo_profile_{slot}.lock"), os.O_RDWR | os.O_CREAT, 0o600)
            try:
                fcntl.flock(fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
            except OSError:
                os.close(fd)
                continue
            try:
                yield os.path.join(base, f"lo_profile_{slot}")
            finally:
                os.close(fd)  # releases the lock
            return
        if time.monotonic() > deadline:
            raise ValueError("The converter is busy. Please try again in a minute.")
        time.sleep(0.2)


def _lock_down_profile(profile: str) -> None:
    user = os.path.join(profile, "user")
    os.makedirs(user, exist_ok=True)
    # Rewritten before every run, so nothing a previous document did to the settings survives.
    with open(os.path.join(user, "registrymodifications.xcu"), "w", encoding="utf-8") as f:
        f.write(_LO_LOCKDOWN)
    # Macros a document might have managed to save into the profile are thrown away too.
    shutil.rmtree(os.path.join(user, "basic"), ignore_errors=True)


def _convert_with_libreoffice(soffice: str, src: str, workdir: str) -> bytes:
    outdir = os.path.join(workdir, "out")
    os.makedirs(outdir, exist_ok=True)
    with _lo_profile() as profile:
        _lock_down_profile(profile)
        cmd = [soffice, f"-env:UserInstallation={Path(profile).resolve().as_uri()}", "--headless", "--invisible",
               "--nologo", "--nodefault", "--nofirststartwizard", "--norestore", "--nolockcheck",
               "--convert-to", "pdf", "--outdir", outdir, src]
        try:
            result = _run(cmd, timeout=180, cwd=workdir)
        except ValueError:
            shutil.rmtree(profile, ignore_errors=True)  # a stuck profile is rebuilt next time
            raise
    pdfs = glob.glob(os.path.join(outdir, "*.pdf"))
    if not pdfs or os.path.getsize(pdfs[0]) == 0:
        detail = (result.stderr or b"").decode("utf-8", "replace").strip().splitlines()
        logging.getLogger(__name__).warning("LibreOffice failed: %s", detail[-1:] if detail else result.returncode)
        raise ValueError("This file could not be converted. It may be damaged, password-protected or not a real Office file.")
    with open(pdfs[0], "rb") as f:
        return f.read()


def warm_up_office() -> None:
    """Builds the LibreOffice profiles in the background so the first real conversion is quick."""
    soffice = find_soffice()
    if not soffice:
        return
    workdir = tempfile.mkdtemp(prefix="ultrapdf_warm_")
    try:
        src = os.path.join(workdir, "warm.txt")
        with open(src, "w") as f:
            f.write("UltraPDF")
        for _ in range(_LO_SLOTS):
            _convert_with_libreoffice(soffice, src, workdir)
    except Exception as exc:
        logging.getLogger(__name__).info("LibreOffice warm-up skipped: %s", exc)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


_OFFICE_EXE = {"Word": "WINWORD.EXE", "Excel": "EXCEL.EXE", "PowerPoint": "POWERPNT.EXE"}


def _office_pids(kind: str) -> set:
    try:
        out = subprocess.run(["tasklist", "/FI", f"IMAGENAME eq {_OFFICE_EXE[kind]}", "/FO", "CSV", "/NH"],
                             capture_output=True, text=True, timeout=15,
                             creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0)).stdout
    except Exception:
        return set()
    return {int(m) for m in re.findall(r'^"[^"]+","(\d+)"', out, re.MULTILINE)}


def _convert_with_ms_office(src: str, out_pdf: str, ext: str) -> bytes:
    """Drives Word/Excel/PowerPoint in a separate process with a time limit, so a stuck
    Office (a hidden dialog, a repair prompt) is stopped instead of hanging forever."""
    kind = "Word" if ext in WORD_EXTS else "Excel" if ext in EXCEL_EXTS else "PowerPoint"
    worker = os.path.join(os.path.dirname(os.path.abspath(__file__)), "office_com_worker.py")
    before = _office_pids(kind)
    try:
        result = _run([sys.executable, worker, kind, src, out_pdf], timeout=180)
    except ValueError:
        for pid in _office_pids(kind) - before:
            subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)
        raise ValueError(f"Microsoft {kind} took too long and was stopped. Installing LibreOffice (free) usually makes this much faster.")
    if result.returncode != 0 or not os.path.isfile(out_pdf) or os.path.getsize(out_pdf) == 0:
        logging.getLogger(__name__).warning("Microsoft %s failed: %s", kind, (result.stderr or b"").decode("utf-8", "replace")[-500:])
        raise ValueError(f"Microsoft {kind} could not convert this file. It may be damaged or password-protected.")
    with open(out_pdf, "rb") as f:
        return f.read()


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", "replace")


_AFTER_PAGESETUP = ("headerFooter", "rowBreaks", "colBreaks", "customProperties", "cellWatches", "ignoredErrors",
                    "smartTags", "drawing", "legacyDrawing", "legacyDrawingHF", "drawingHF", "picture", "oleObjects",
                    "controls", "webPublishItems", "tableParts", "extLst")


def _col_number(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + ord(ch) - 64
    return n


def _fit_sheet_xml(xml: str) -> str:
    """Sets one worksheet to print all columns on one page wide (landscape when wide)."""
    m = re.search(r'<(?:\w+:)?dimension\s+ref="[A-Z]+\d+:([A-Z]+)\d+"', xml)
    wide = bool(m) and _col_number(m.group(1)) > 7
    # <sheetPr><pageSetUpPr fitToPage="1"/></sheetPr>
    if re.search(r"<pageSetUpPr\b", xml):
        xml = re.sub(r"<pageSetUpPr\b([^>]*?)(/?)>", lambda mm: "<pageSetUpPr" + re.sub(r'\s+fitToPage="[^"]*"', "", mm.group(1)) + ' fitToPage="1"' + mm.group(2) + ">", xml, count=1)
    elif re.search(r"<sheetPr\b[^>]*/>", xml):
        xml = re.sub(r"<sheetPr\b([^>]*)/>", r'<sheetPr\1><pageSetUpPr fitToPage="1"/></sheetPr>', xml, count=1)
    elif re.search(r"<sheetPr\b", xml):
        xml = xml.replace("</sheetPr>", '<pageSetUpPr fitToPage="1"/></sheetPr>', 1)
    else:
        xml = re.sub(r"(<worksheet\b[^>]*>)", r'\1<sheetPr><pageSetUpPr fitToPage="1"/></sheetPr>', xml, count=1)
    attrs = ' fitToWidth="1" fitToHeight="0"' + (' orientation="landscape"' if wide else "")
    if re.search(r"<pageSetup\b", xml):
        def fix(mm):
            a = re.sub(r'\s+(fitToWidth|fitToHeight|scale' + ("|orientation" if wide else "") + r')="[^"]*"', "", mm.group(1))
            return "<pageSetup" + a + attrs + mm.group(2) + ">"
        xml = re.sub(r"<pageSetup\b([^>]*?)(/?)>", fix, xml, count=1)
    else:
        tag = "<pageSetup" + attrs + "/>"
        if not re.search(r"<pageMargins\b", xml):
            tag = '<pageMargins left="0.5" right="0.5" top="0.6" bottom="0.6" header="0.3" footer="0.3"/>' + tag
            where = None
        else:
            where = re.search(r"<pageMargins\b[^>]*/>", xml)
        if where:
            xml = xml[:where.end()] + tag + xml[where.end():]
        else:
            nxt = re.search(r"<(?:" + "|".join(_AFTER_PAGESETUP) + r")\b", xml)
            pos = nxt.start() if nxt else xml.rindex("</worksheet>")
            xml = xml[:pos] + tag + xml[pos:]
    return xml


MAX_UNZIPPED = 600 * 1024 * 1024
MAX_ZIP_ENTRIES = 20000


def check_zip_container(data: bytes) -> None:
    """Office (docx/xlsx/pptx) and OpenDocument files are zip archives. Refuses "zip bombs":
    small uploads that expand to gigabytes, or that hold huge numbers of parts."""
    if data[:4] != b"PK\x03\x04":
        return  # older binary formats (.doc/.xls/.ppt) or not a zip; the converter decides
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        raise ValueError("This file is damaged or not a real Office document.")
    if len(infos) > MAX_ZIP_ENTRIES:
        raise ValueError("This document has too many parts to convert safely.")
    total = sum(i.file_size for i in infos)
    if total > MAX_UNZIPPED:
        raise ValueError("This document expands to more than 600 MB when opened, so it can't be converted safely.")
    for i in infos:
        if i.compress_size and i.file_size / i.compress_size > 1000 and i.file_size > 50 * 1024 * 1024:
            raise ValueError("This document looks like a 'zip bomb' and can't be converted safely.")


def fit_xlsx_columns(data: bytes) -> bytes:
    """Returns the workbook with every sheet set to fit all its columns on one page width.
    Edits only the page-setup XML, so formatting, charts and images are kept."""
    try:
        zin = zipfile.ZipFile(io.BytesIO(data))
        names = zin.namelist()
    except zipfile.BadZipFile:
        return data
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            body = zin.read(info.filename)
            if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                try:
                    body = _fit_sheet_xml(body.decode("utf-8")).encode("utf-8")
                except Exception:
                    pass
            zout.writestr(info, body)
    return out.getvalue() if names else data


def office_to_pdf(data: bytes, filename: str) -> bytes:
    """Converts .doc/.docx/.odt/.rtf/.txt/.xls/.xlsx/.ods/.csv/.ppt/.pptx/.odp to PDF."""
    if not data:
        raise ValueError("The file is empty.")
    ext = Path(filename or "").suffix.lower()
    if ext not in OFFICE_EXTS:
        raise ValueError(f"Unsupported file type '{ext or 'none'}'. Use one of: " + ", ".join(sorted(OFFICE_EXTS)) + ".")
    if ext == ".csv":
        return csv_to_pdf(data)
    soffice = find_soffice()
    ms_apps = _ms_office_apps() if not soffice else {}
    kind = "Word" if ext in WORD_EXTS else "Excel" if ext in EXCEL_EXTS else "PowerPoint"
    if not soffice and not ms_apps.get(kind):
        # Plain text and Excel files can still be drawn without an office suite.
        if ext == ".txt":
            return text_to_pdf(_decode_text(data))
        if ext in (".xlsx", ".xlsm"):
            check_zip_container(data)
            return xlsx_to_pdf_simple(fit_xlsx_columns(data))
        raise ToolUnavailableError("Office conversion needs LibreOffice (free) or Microsoft Office on this computer.")
    check_zip_container(data)
    if ext in (".xlsx", ".xlsm"):
        data = fit_xlsx_columns(data)
    workdir = tempfile.mkdtemp(prefix="ultrapdf_office_")
    try:
        src = os.path.join(workdir, _safe_name(filename, ext))
        with open(src, "wb") as f:
            f.write(data)
        if soffice:
            pdf = _convert_with_libreoffice(soffice, src, workdir)
        else:
            pdf = _convert_with_ms_office(os.path.abspath(src), os.path.abspath(os.path.join(workdir, "out.pdf")), ext)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return pdf


# ---------------------------------------------------------------------------
# 2. Markdown / HTML / CSV / text -> PDF  (PyMuPDF Story)
# ---------------------------------------------------------------------------

BASE_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.4; color: #222222; }
h1 { font-size: 22pt; margin: 0 0 10pt 0; color: #111111; }
h2 { font-size: 17pt; margin: 14pt 0 8pt 0; color: #111111; }
h3 { font-size: 14pt; margin: 12pt 0 6pt 0; }
h4, h5, h6 { font-size: 12pt; margin: 10pt 0 4pt 0; }
p { margin: 0 0 8pt 0; }
a { color: #1a5fb4; text-decoration: underline; }
ul, ol { margin: 0 0 8pt 0; }
li { margin: 0 0 2pt 0; }
code { font-family: monospace; font-size: 9.5pt; background-color: #f2f2f2; }
pre { font-family: monospace; font-size: 9pt; background-color: #f4f4f4; border: 1px solid #dddddd;
      padding: 6pt; margin: 0 0 8pt 0; white-space: pre-wrap; }
pre code { background-color: #f4f4f4; }
blockquote { margin: 0 0 8pt 12pt; padding-left: 8pt; border-left: 3px solid #cccccc; color: #555555; }
table { border-collapse: collapse; margin: 0 0 10pt 0; }
th, td { border: 1px solid #999999; padding: 3pt 5pt; vertical-align: top; text-align: left; }
th { font-weight: bold; background-color: #eeeeee; }
hr { border: 0; border-top: 1px solid #cccccc; }
img { max-width: 100%; }
"""


def _story_to_pdf(html: str, css: str = BASE_CSS, landscape: bool = False) -> bytes:
    """Lays HTML out over as many A4 pages as needed and returns the PDF."""
    mediabox = pymupdf.Rect(0, 0, A4.height, A4.width) if landscape else pymupdf.Rect(A4)
    where = mediabox + (MARGIN, MARGIN, -MARGIN, -MARGIN)
    try:
        story = pymupdf.Story(html=html, user_css=css)
    except Exception:
        raise ValueError("The document could not be laid out. Check that it is valid HTML.")

    def rectfn(rect_num, _filled):
        if rect_num >= MAX_PAGES_OUT:
            raise ValueError(f"The document is too long (over {MAX_PAGES_OUT} pages).")
        return mediabox, where, None

    doc = story.write_with_links(rectfn)  # keeps <a href> links clickable
    try:
        return doc.tobytes(garbage=3, deflate=True)
    finally:
        doc.close()


_DANGEROUS_BLOCKS = re.compile(r"<(script|noscript|iframe|object|embed|applet|frameset|frame|template|svg|math|form|video|audio)\b.*?</\1\s*>",
                               re.IGNORECASE | re.DOTALL)
_DANGEROUS_SINGLE = re.compile(r"<(script|iframe|object|embed|link|meta|base|frame|input|source|track)\b[^>]*>", re.IGNORECASE)
_EVENT_ATTR = re.compile(r"""\s+on[a-z]+\s*=\s*("[^"]*"|'[^']*'|[^\s>]+)""", re.IGNORECASE)
_REMOTE_ATTR = re.compile(r"""\s+(src|href|background|poster|srcset)\s*=\s*("[^"]*"|'[^']*'|[^\s>]+)""", re.IGNORECASE)
_CSS_URL = re.compile(r"url\s*\([^)]*\)", re.IGNORECASE)
_CSS_IMPORT = re.compile(r"@import[^;]*;?", re.IGNORECASE)


def _sanitize_html(html: str) -> str:
    """Removes scripts and anything that could make the renderer load another file or URL."""
    html = re.sub(r"<!--.*?-->", "", html, flags=re.DOTALL)
    html = _DANGEROUS_BLOCKS.sub("", html)
    html = _DANGEROUS_SINGLE.sub("", html)
    html = re.sub(r"</(script|iframe|object|embed)\s*>", "", html, flags=re.IGNORECASE)
    html = _EVENT_ATTR.sub("", html)

    def fix_attr(m: "re.Match") -> str:
        name = m.group(1).lower()
        value = m.group(2).strip("\"'").strip()
        if name == "href" and re.match(r"(https?:|mailto:|#)", value, re.IGNORECASE):
            return m.group(0)  # a clickable link is fine; nothing is fetched
        if name == "src" and value.lower().startswith("data:image/"):
            return m.group(0)  # inline image, no fetch
        return ""

    html = _REMOTE_ATTR.sub(fix_attr, html)
    html = _CSS_IMPORT.sub("", html)
    html = _CSS_URL.sub("none", html)
    return html


def html_to_pdf(html: str) -> bytes:
    """Renders an HTML page to PDF. Scripts are dropped and nothing is fetched from the network."""
    if not html or not html.strip():
        raise ValueError("The HTML is empty.")
    if len(html) > 20_000_000:
        raise ValueError("The HTML is too large (over 20 MB).")
    return _story_to_pdf(_sanitize_html(html))


# ----- Markdown -----

def _md_inline(text: str) -> str:
    """Inline Markdown (code, links, bold, italic) for the built-in fallback converter."""
    codes: List[str] = []

    def keep_code(m: "re.Match") -> str:
        codes.append(f"<code>{html_lib.escape(m.group(1))}</code>")
        return f"\x00{len(codes) - 1}\x00"

    text = re.sub(r"`([^`]+)`", keep_code, text)
    text = html_lib.escape(text, quote=False)
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)\s]+)[^)]*\)", lambda m: f'<a href="{html_lib.escape(m.group(2))}">{m.group(1)}</a>', text)
    text = re.sub(r"(\*\*|__)(.+?)\1", r"<b>\2</b>", text)
    text = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"<i>\1</i>", text)
    text = re.sub(r"(?<![\w_])_(?!\s)(.+?)(?<!\s)_(?![\w_])", r"<i>\1</i>", text)
    text = re.sub(r"~~(.+?)~~", r"<s>\1</s>", text)
    return re.sub(r"\x00(\d+)\x00", lambda m: codes[int(m.group(1))], text)


def _simple_markdown(md: str) -> str:
    """A small Markdown converter used when the 'markdown' package is not installed."""
    lines = md.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    out: List[str] = []
    para: List[str] = []
    i = 0

    def flush_para():
        if para:
            out.append("<p>" + _md_inline(" ".join(para)) + "</p>")
            para.clear()

    def split_row(row: str) -> List[str]:
        row = row.strip()
        if row.startswith("|"):
            row = row[1:]
        if row.endswith("|"):
            row = row[:-1]
        return [c.strip() for c in row.split("|")]

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        fence = re.match(r"^(```|~~~)", stripped)
        if fence:
            flush_para()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(fence.group(1)):
                code.append(lines[i])
                i += 1
            out.append("<pre><code>" + html_lib.escape("\n".join(code)) + "</code></pre>")
            i += 1
            continue
        if not stripped:
            flush_para()
            i += 1
            continue
        h = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", stripped)
        if h:
            flush_para()
            n = len(h.group(1))
            out.append(f"<h{n}>{_md_inline(h.group(2))}</h{n}>")
            i += 1
            continue
        if re.match(r"^([-*_])(\s*\1){2,}\s*$", stripped):
            flush_para()
            out.append("<hr/>")
            i += 1
            continue
        if "|" in stripped and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$", lines[i + 1]):
            flush_para()
            head = split_row(stripped)
            rows = []
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_row(lines[i]))
                i += 1
            t = ["<table><thead><tr>" + "".join(f"<th>{_md_inline(c)}</th>" for c in head) + "</tr></thead><tbody>"]
            for r in rows:
                r = (r + [""] * len(head))[:len(head)]
                t.append("<tr>" + "".join(f"<td>{_md_inline(c)}</td>" for c in r) + "</tr>")
            out.append("".join(t) + "</tbody></table>")
            continue
        if stripped.startswith(">"):
            flush_para()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].lstrip())
                i += 1
            out.append("<blockquote>" + _simple_markdown("\n".join(quote)) + "</blockquote>")
            continue
        lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if lm:
            flush_para()
            ordered = lm.group(2)[0].isdigit()
            tag = "ol" if ordered else "ul"
            items = []
            while i < len(lines):
                m2 = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", lines[i])
                if m2 and m2.group(2)[0].isdigit() == ordered:
                    items.append(m2.group(3))
                elif lines[i].strip() and lines[i].startswith((" ", "\t")) and items:
                    items[-1] += " " + lines[i].strip()
                else:
                    break
                i += 1
            out.append(f"<{tag}>" + "".join(f"<li>{_md_inline(it)}</li>" for it in items) + f"</{tag}>")
            continue
        if re.match(r"^ {4}|\t", line):
            flush_para()
            code = []
            while i < len(lines) and (re.match(r"^ {4}|\t", lines[i]) or not lines[i].strip()):
                code.append(re.sub(r"^ {4}|\t", "", lines[i], count=1))
                i += 1
            out.append("<pre><code>" + html_lib.escape("\n".join(code).rstrip("\n")) + "</code></pre>")
            continue
        para.append(stripped)
        i += 1
    flush_para()
    return "\n".join(out)


def markdown_to_html(text: str) -> str:
    try:
        import markdown as md_lib
        return md_lib.markdown(text, extensions=["tables", "fenced_code", "sane_lists"], output_format="html")
    except ImportError:
        return _simple_markdown(text)


def markdown_to_pdf(text: str) -> bytes:
    """Renders Markdown (headings, lists, tables, code, links) to a clean A4 PDF."""
    if not text or not text.strip():
        raise ValueError("The Markdown text is empty.")
    if len(text) > 20_000_000:
        raise ValueError("The Markdown text is too large (over 20 MB).")
    return _story_to_pdf(_sanitize_html(markdown_to_html(text)))


def text_to_pdf(text: str) -> bytes:
    """Plain text to PDF, keeping line breaks and wrapping long lines."""
    if not text or not text.strip():
        raise ValueError("The text file is empty.")
    paras = "".join(f"<p>{html_lib.escape(line) or '&#160;'}</p>" for line in text.replace("\r\n", "\n").split("\n"))
    css = BASE_CSS + "p { font-family: monospace; font-size: 10pt; margin: 0; }"
    return _story_to_pdf(paras, css=css)


# ----- CSV -----

def _read_csv(data: bytes) -> List[List[str]]:
    if not data or not data.strip():
        raise ValueError("The CSV file is empty.")
    text = _decode_text(data)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        if "\t" in sample.split("\n", 1)[0] and "," not in sample.split("\n", 1)[0]:
            dialect = csv.excel_tab
    try:
        rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    except csv.Error:
        raise ValueError("This CSV file could not be read.")
    if not rows:
        raise ValueError("The CSV file has no data.")
    return rows


def _display(value: Any) -> str:
    """How a cell reads on paper: long floats rounded like Excel's General format."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return "%.10g" % value
    if hasattr(value, "isoformat"):
        text = value.isoformat(sep=" ") if hasattr(value, "date") and hasattr(value, "hour") else value.isoformat()
        return text[:10] if text.endswith(" 00:00:00") else text
    text = str(value)
    if re.fullmatch(r"-?\d+\.\d{11,}", text.strip()):
        try:
            return "%.10g" % float(text)
        except ValueError:
            pass
    return text


def _wrap(text: str, width: float, size: float) -> List[str]:
    lines: List[str] = []
    for para in text.split("\n"):
        line = ""
        for word in para.split(" "):
            trial = f"{line} {word}" if line else word
            if not line or pymupdf.get_text_length(trial, "helv", size) <= width:
                line = trial
            else:
                lines.append(line)
                line = word
        lines.append(line)
    return lines or [""]


def table_pdf(sheets: List[Tuple[str, List[List[Any]]]]) -> bytes:
    """Draws tables so every column fits across the page: landscape, then smaller text,
    then a wider page. The header row repeats on every page."""
    margin, base, pad = 28.0, 10.0, 3.0
    doc = pymupdf.open()
    for title, raw in sheets:
        rows = [[_display(c) for c in r] for r in raw]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        # Width each column needs at 10pt, with very long cells wrapped at 240pt.
        natural = []
        for c in range(ncols):
            w = max(pymupdf.get_text_length(max(r[c].split("\n"), key=len), "hebo" if i == 0 else "helv", base)
                    for i, r in enumerate(rows[:2000]))
            natural.append(min(max(w, 12.0), 240.0) + 2 * pad)
        total = sum(natural)
        page_w, page_h = A4.height, A4.width  # landscape A4
        if ncols <= 5 and total <= A4.width - 2 * margin:
            page_w, page_h = A4.width, A4.height
        scale = min(1.0, (page_w - 2 * margin) / total)
        if base * scale < 6:
            a3 = pymupdf.paper_rect("a3")
            page_w, page_h = a3.height, a3.width
            scale = min(1.0, (page_w - 2 * margin) / total)
        if base * scale < 5:
            scale = 0.5  # 5pt text on a page as wide as the table needs
            page_w = min(total * scale + 2 * margin, 14000.0)
            scale = min(scale, (page_w - 2 * margin) / total)
        size = base * scale
        widths = [w * scale for w in natural]
        line_h = size * 1.25
        p = pad * scale

        def row_lines(r):
            return [_wrap(cell, w - 2 * p, size) for cell, w in zip(r, widths)]

        header = rows[0]
        header_lines = row_lines(header)
        page = None
        y = 0.0

        def new_page():
            nonlocal page, y
            if len(doc) >= MAX_PAGES_OUT:
                raise ValueError(f"The table is too long (over {MAX_PAGES_OUT} pages).")
            page = doc.new_page(width=page_w, height=page_h)
            y = margin
            if title:
                page.insert_text((margin, y + 10), title, fontname="hebo", fontsize=11, color=(0.2, 0.2, 0.2))
                y += 18
            draw(header_lines, bold=True)

        def draw(lines, bold=False):
            nonlocal y
            h = max(len(l) for l in lines) * line_h + 2 * p
            x = margin
            for cell_lines, w in zip(lines, widths):
                rect = pymupdf.Rect(x, y, x + w, y + h)
                page.draw_rect(rect, color=(0.6, 0.6, 0.6), fill=(0.93, 0.93, 0.93) if bold else None, width=0.4)
                ty = y + p + size
                for line in cell_lines:
                    page.insert_text((x + p, ty), line, fontname="hebo" if bold else "helv", fontsize=size, color=(0.1, 0.1, 0.1))
                    ty += line_h
                x += w
            y += h

        new_page()
        for r in rows[1:]:
            lines = row_lines(r)
            h = max(len(l) for l in lines) * line_h + 2 * p
            if y + h > page_h - margin:
                new_page()
            draw(lines)
    if len(doc) == 0:
        raise ValueError("There is no data to convert.")
    try:
        return doc.tobytes(garbage=3, deflate=True)
    finally:
        doc.close()


def csv_to_pdf(data: bytes) -> bytes:
    """Renders a CSV file as a table with every column on the page."""
    rows = _read_csv(data)
    if len(rows) > 100_000:
        raise ValueError("This CSV has more than 100,000 rows. Split it into smaller files first.")
    return table_pdf([("", rows)])


def xlsx_to_pdf_simple(data: bytes) -> bytes:
    """Excel workbook to PDF without an office suite: values only, every sheet, all columns."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise ValueError("This Excel file could not be read. It may be damaged or password-protected.")
    sheets = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        sheets.append((ws.title if len(wb.worksheets) > 1 else "", rows))
    wb.close()
    return table_pdf(sheets)


# ---------------------------------------------------------------------------
# 3. PDF -> Excel
# ---------------------------------------------------------------------------

_NUM = re.compile(r"^(?P<neg>-|\()?\s*(?P<cur>[$€£¥])?\s*(?P<num>\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?|\.\d+)\s*(?P<pct>%)?\)?$")


def _to_number(value: str):
    """'1,234.50' -> 1234.5, '12%' -> 0.12, '(5)' -> -5; anything else stays text."""
    s = value.strip()
    m = _NUM.match(s)
    if not m:
        return value
    num = m.group("num").replace(",", "")
    digits = num.replace(".", "")
    if (len(num) > 1 and num.startswith("0") and not num.startswith("0.")) or len(digits) > 15:
        return value  # IDs, zip codes, account numbers
    if m.group("neg") == "(" and not s.endswith(")"):
        return value
    n: Any = float(num) if "." in num else int(num)
    if m.group("neg"):
        n = -n
    if m.group("pct"):
        n = n / 100
    return n


def _tables_pymupdf(doc) -> Dict[int, List[List[List[Any]]]]:
    found: Dict[int, List[List[List[Any]]]] = {}
    for page in doc:
        try:
            tabs = page.find_tables()
        except Exception:
            continue
        for t in tabs.tables:
            rows = t.extract()
            if rows and any(any(c for c in r) for r in rows):
                found.setdefault(page.number, []).append(rows)
    return found


def _tables_pdfplumber(pdf: bytes) -> Dict[int, List[List[List[Any]]]]:
    found: Dict[int, List[List[List[Any]]]] = {}
    try:
        import pdfplumber
    except ImportError:
        return found
    try:
        with pdfplumber.open(io.BytesIO(pdf)) as p:
            for i, page in enumerate(p.pages):
                for rows in page.extract_tables() or []:
                    if rows and any(any(c for c in r) for r in rows):
                        found.setdefault(i, []).append(rows)
    except Exception:
        pass
    return found


def _autosize(ws) -> None:
    widths: Dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                longest = max(len(line) for line in str(c.value).split("\n"))
                widths[c.column] = max(widths.get(c.column, 0), longest)
    from openpyxl.utils import get_column_letter
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(8, w + 2), 80)


def pdf_to_excel(pdf: bytes) -> bytes:
    """Extracts tables into an .xlsx workbook: one sheet per page with tables, numbers as numbers.
    A PDF without any tables gives one sheet holding its text lines."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        raise ToolUnavailableError("PDF to Excel needs the 'openpyxl' package on the server.")
    doc = _open_pdf(pdf)
    try:
        tables = _tables_pymupdf(doc)
        if not tables:
            tables = _tables_pdfplumber(pdf)
        wb = Workbook()
        wb.remove(wb.active)
        bold = Font(bold=True)
        if tables:
            for pno in sorted(tables):
                ws = wb.create_sheet(f"Page {pno + 1}")
                r = 1
                for rows in tables[pno]:
                    header_row = r
                    for row in rows:
                        for ci, val in enumerate(row, start=1):
                            if val is None or val == "":
                                continue
                            text = _xml_safe(str(val))
                            cell = ws.cell(row=r, column=ci, value=_to_number(text) if r != header_row else text)
                            if r == header_row:
                                cell.font = bold
                            if "\n" in text:
                                cell.alignment = Alignment(wrap_text=True, vertical="top")
                        r += 1
                    r += 1  # blank row between tables
                _autosize(ws)
        else:
            ws = wb.create_sheet("Text")
            ws.append(["Page", "Text"])
            ws["A1"].font = bold
            ws["B1"].font = bold
            for page in doc:
                for line in page.get_text().splitlines():
                    if line.strip():
                        ws.append([page.number + 1, _xml_safe(line)])
            if ws.max_row == 1:
                ws.append([None, "No tables or text were found in this PDF (it may be a scan - run OCR first)."])
            _autosize(ws)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    finally:
        doc.close()


# ---------------------------------------------------------------------------
# 4. PDF -> PowerPoint
# ---------------------------------------------------------------------------

EMU_PER_PT = 12700
SLIDE_MIN, SLIDE_MAX = 914400, 51206400  # PowerPoint's limits: 1 inch to 56 inches


def pdf_to_pptx(pdf: bytes) -> bytes:
    """One slide per page: the page as a full-slide picture, its text in the speaker notes."""
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        raise ToolUnavailableError("PDF to PowerPoint needs the 'python-pptx' package on the server.")
    doc = _open_pdf(pdf)
    try:
        first = doc[0].rect
        width = 9144000  # 10 inches, like a standard 16:9 / 4:3 slide
        height = int(width * first.height / first.width)
        if height > SLIDE_MAX:
            width, height = int(SLIDE_MAX * first.width / first.height), SLIDE_MAX
        width = max(SLIDE_MIN, min(SLIDE_MAX, width))
        height = max(SLIDE_MIN, min(SLIDE_MAX, height))
        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(width), Emu(height)
        blank = prs.slide_layouts[6]
        for page in doc:
            r = page.rect
            dpi = 150
            longest_in = max(r.width, r.height) / 72
            if longest_in * dpi > 4000:
                dpi = max(36, int(4000 / longest_in))
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            img = pix.tobytes("jpeg", jpg_quality=88)
            slide = prs.slides.add_slide(blank)
            # Fit the page inside the slide (pages may differ in shape from the first one).
            scale = min(width / r.width, height / r.height)
            pw, ph = min(width, round(r.width * scale)), min(height, round(r.height * scale))
            slide.shapes.add_picture(io.BytesIO(img), Emu((width - pw) // 2), Emu((height - ph) // 2), Emu(pw), Emu(ph))
            text = _xml_safe(page.get_text().strip())
            if text:
                slide.notes_slide.notes_text_frame.text = text
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()
    finally:
        doc.close()


# ---------------------------------------------------------------------------
# 5. PDF -> Word (layout preserving)
# ---------------------------------------------------------------------------

def pdf_to_docx(pdf: bytes) -> bytes:
    """Layout-preserving PDF to .docx using pdf2docx. Raises ToolUnavailableError when
    pdf2docx isn't installed, so the caller can fall back to the text-only converter."""
    try:
        logging.getLogger("pdf2docx").setLevel(logging.ERROR)
        from pdf2docx import Converter
    except ImportError:
        raise ToolUnavailableError("Layout-preserving Word conversion needs the 'pdf2docx' package on the server.")
    _open_pdf(pdf).close()  # friendly errors for broken / locked files
    workdir = tempfile.mkdtemp(prefix="ultrapdf_docx_")
    try:
        src = os.path.join(workdir, "in.pdf")
        dst = os.path.join(workdir, "out.docx")
        with open(src, "wb") as f:
            f.write(pdf)
        previous = logging.root.manager.disable
        logging.disable(logging.INFO)  # pdf2docx logs every step at INFO level
        try:
            cv = Converter(src)
            try:
                cv.convert(dst, multi_processing=False)
            finally:
                cv.close()
        except Exception as exc:
            logging.getLogger(__name__).warning("pdf2docx failed: %s", exc)
            raise ValueError("This PDF's layout could not be converted to Word.")
        finally:
            logging.disable(previous)
        if not os.path.isfile(dst) or os.path.getsize(dst) == 0:
            raise ValueError("This PDF's layout could not be converted to Word.")
        with open(dst, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


# ---------------------------------------------------------------------------
# 6. Repair
# ---------------------------------------------------------------------------

def _page_count(pdf: bytes) -> int:
    try:
        d = pymupdf.open(stream=pdf, filetype="pdf")
        n = d.page_count if not d.needs_pass else -1
        d.close()
        return n
    except Exception:
        return 0


def _repair_pikepdf(data: bytes) -> Tuple[Optional[bytes], List[str]]:
    try:
        import pikepdf
    except ImportError:
        return None, []
    try:
        with pikepdf.open(io.BytesIO(data), attempt_recovery=True, suppress_warnings=True) as p:
            warnings = list(p.get_warnings())
            if len(p.pages) == 0:
                return None, []
            buf = io.BytesIO()
            p.save(buf, compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate)
            return buf.getvalue(), warnings
    except Exception:
        return None, []


def _repair_pymupdf(data: bytes) -> Tuple[Optional[bytes], bool]:
    try:
        d = pymupdf.open(stream=data, filetype="pdf")
        if d.needs_pass or d.page_count == 0:
            return None, False
        repaired = bool(getattr(d, "is_repaired", False))
        out = d.tobytes(garbage=3, clean=True, deflate=True)
        d.close()
        return out, repaired
    except Exception:
        return None, False


def repair_pdf(pdf: bytes) -> Dict[str, Any]:
    """Rebuilds a damaged PDF. Returns {"bytes": fixed PDF, "report": [what was done]}."""
    if not pdf:
        raise ValueError("The file is empty.")
    report: List[str] = []
    start = pdf.find(b"%PDF-")
    if start < 0:
        raise ValueError("This file does not contain any PDF data, so it cannot be repaired.")
    data = pdf
    if start > 0:
        data = pdf[start:]
        report.append(f"Removed {start} bytes of junk in front of the PDF header.")
    try:
        locked = pymupdf.open(stream=data, filetype="pdf").needs_pass
    except Exception:
        locked = False
    if locked:
        raise ValueError("This PDF is password-protected. Unlock it first, then repair it.")

    original_pages = 0
    try:
        import pikepdf
        with pikepdf.open(io.BytesIO(data), attempt_recovery=False) as p:
            original_pages = len(p.pages)
        healthy = True
    except Exception:
        healthy = False

    candidates: List[Tuple[int, bytes, List[str]]] = []
    out, warnings = _repair_pikepdf(data)
    if out:
        notes = []
        text = " ".join(warnings).lower()
        if "xref" in text or "reconstruct" in text or "startxref" in text:
            notes.append("Rebuilt the damaged cross-reference table (the file's index of objects).")
        if "eof" in text or "truncat" in text or "end of file" in text:
            notes.append("Recovered content from a file that was cut off before the end.")
        if "stream" in text and "length" in text:
            notes.append("Fixed streams with wrong length information.")
        if warnings and not notes:
            notes.append(f"Fixed {len(warnings)} structural problem(s) found while reading the file.")
        candidates.append((_page_count(out), out, notes))
    out2, repaired = _repair_pymupdf(data)
    if out2:
        notes = ["Rebuilt the file structure from its objects."] if repaired else []
        candidates.append((_page_count(out2), out2, notes))

    if not candidates or max(c[0] for c in candidates) <= 0:
        # Salvage: cut off anything after the last end-of-file marker, or try each PDF header.
        tries = []
        eof = data.rfind(b"%%EOF")
        if eof > 0 and data[eof + 5:].strip():
            tries.append((data[:eof + 5] + b"\n", "Cut off junk data after the end of the file."))
        for m in re.finditer(rb"%PDF-", data):
            if m.start() > 0:
                tries.append((data[m.start():], "Recovered an embedded copy of the PDF found inside the file."))
        for blob, note in tries:
            for fn in (_repair_pikepdf, _repair_pymupdf):
                res, _ = fn(blob)
                if res and _page_count(res) > 0:
                    candidates.append((_page_count(res), res, [note]))
                    break
            if candidates and max(c[0] for c in candidates) > 0:
                break

    candidates = [c for c in candidates if c[0] > 0]
    if not candidates:
        raise ValueError("This PDF is too badly damaged; no pages could be recovered.")
    pages, best, notes = max(candidates, key=lambda c: (c[0], -len(c[1])))
    report += notes
    if healthy and not report:
        report.append("No damage was found. The file was re-saved cleanly with unused objects removed.")
    elif not notes and not healthy:
        report.append("Rebuilt the damaged file structure.")
    if original_pages and pages < original_pages:
        report.append(f"Only {pages} of {original_pages} pages could be recovered.")
    report.append(f"The repaired PDF has {pages} page{'s' if pages != 1 else ''}.")
    return {"bytes": best, "report": report}


# ---------------------------------------------------------------------------
# 7. Privacy scan / strip
# ---------------------------------------------------------------------------

_EMAIL_RE = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}\b")
_PHONE_RE = re.compile(r"(?<![\w+])(?<!\d[\s.-])(?:\+\d{1,3}[\s.-]?)?(?:\(\d{2,4}\)[\s.-]?|\d{2,4}[\s.-])\d{3,4}[\s.-]\d{3,4}(?![\w])(?![\s.-]?\d)"
                       r"|(?<![\w+])\+\d{8,14}(?!\w)")
_CARD_RE = re.compile(r"(?<!\d)(?:\d[ -]?){12,18}\d(?!\d)")
_IBAN_RE = re.compile(r"\b[A-Z]{2}\d{2}(?:[ ]?[A-Z0-9]{4}){2,7}(?:[ ]?[A-Z0-9]{1,4})?\b")
_ACTION_TYPES = {"/JavaScript": "JavaScript", "/Launch": "launch-a-program", "/SubmitForm": "submit-form",
                 "/ImportData": "import-data", "/GoToR": "open-another-file", "/GoToE": "open-embedded-file"}
SEVERITY_PENALTY = {"high": 25, "medium": 10, "low": 3}


def _luhn(digits: str) -> bool:
    total = 0
    for i, ch in enumerate(reversed(digits)):
        d = int(ch)
        if i % 2:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def _iban_ok(raw: str) -> bool:
    s = raw.replace(" ", "")
    if not 15 <= len(s) <= 34:
        return False
    s = s[4:] + s[:4]
    try:
        return int("".join(str(int(c, 36)) for c in s)) % 97 == 1
    except ValueError:
        return False


def _find_sensitive(text: str) -> Dict[str, List[str]]:
    found: Dict[str, List[str]] = {"email": [], "phone": [], "card": [], "iban": []}
    for m in _EMAIL_RE.finditer(text):
        found["email"].append(m.group(0))
    card_spans = []
    for m in _CARD_RE.finditer(text):
        digits = re.sub(r"\D", "", m.group(0))
        if 13 <= len(digits) <= 19 and _luhn(digits) and len(set(digits)) > 1:
            found["card"].append(m.group(0).strip())
            card_spans.append(m.span())
    for m in _IBAN_RE.finditer(text):
        if _iban_ok(m.group(0)):
            found["iban"].append(m.group(0))
            card_spans.append(m.span())
    for m in _PHONE_RE.finditer(text):
        if any(a <= m.start() < b for a, b in card_spans):
            continue
        digits = re.sub(r"\D", "", m.group(0))
        if 7 <= len(digits) <= 15:
            found["phone"].append(m.group(0).strip())
    return {k: list(dict.fromkeys(v)) for k, v in found.items()}


def _actions_in(pdf: bytes) -> Dict[str, Any]:
    """Walks every object with pikepdf looking for scripts, auto-run and launch actions."""
    info: Dict[str, Any] = {"actions": {}, "open_action": False, "uris": set(), "incremental": 0, "xmp": False}
    try:
        import pikepdf
    except ImportError:
        return info
    try:
        p = pikepdf.open(io.BytesIO(pdf))
    except Exception:
        return info
    with p:
        root = p.Root
        if "/OpenAction" in root:
            oa = root.OpenAction
            # A plain "open at page 1" destination is harmless; scripts and launches are not.
            if isinstance(oa, pikepdf.Dictionary) and str(oa.get("/S", "")) in _ACTION_TYPES:
                info["open_action"] = True
        if "/AA" in root:
            info["open_action"] = True
        names = root.get("/Names")
        if isinstance(names, pikepdf.Dictionary) and "/JavaScript" in names:
            info["actions"]["JavaScript"] = info["actions"].get("JavaScript", 0) + 1
        info["xmp"] = "/Metadata" in root
        for obj in p.objects:
            if not isinstance(obj, pikepdf.Dictionary):
                continue
            s = obj.get("/S")
            if s is not None and str(s) in _ACTION_TYPES and ("/JS" in obj or str(s) != "/JavaScript"):
                label = _ACTION_TYPES[str(s)]
                info["actions"][label] = info["actions"].get(label, 0) + 1
            for key in ("/A", "/AA"):
                sub = obj.get(key)
                if isinstance(sub, pikepdf.Dictionary) and not sub.is_indirect:
                    subs = [sub] if key == "/A" else [v for v in sub.values() if isinstance(v, pikepdf.Dictionary) and not v.is_indirect]
                    for a in subs:
                        st = str(a.get("/S", ""))
                        if st in _ACTION_TYPES:
                            info["actions"][_ACTION_TYPES[st]] = info["actions"].get(_ACTION_TYPES[st], 0) + 1
                        if st == "/URI" and "/URI" in a:
                            info["uris"].add(str(a.URI))
            if str(obj.get("/S", "")) == "/URI" and "/URI" in obj:
                info["uris"].add(str(obj.URI))
        linearized = p.is_linearized
    eofs = len(re.findall(rb"%%EOF", pdf))
    info["incremental"] = max(0, eofs - (2 if linearized else 1))
    return info


def _hidden_text(doc) -> Dict[str, List[str]]:
    """Text a reader can't see: invisible render mode, white-on-white, tiny, or off the page."""
    kinds: Dict[str, List[str]] = {"invisible": [], "white": [], "tiny": [], "offpage": []}
    for page in doc:
        if not hasattr(page, "get_texttrace"):
            break
        try:
            trace = page.get_texttrace()
        except Exception:
            continue
        has_images = bool(page.get_images())
        prect = page.rect
        for span in trace:
            text = "".join(chr(c[0]) for c in span.get("chars", ()) if c[0] > 0).strip()
            if not text:
                continue
            bbox = pymupdf.Rect(span["bbox"])
            if span.get("type") == 3 or span.get("opacity", 1) == 0:
                # An invisible layer over a scan is normal OCR output; only flag it on text pages.
                if not has_images:
                    kinds["invisible"].append(text)
            elif tuple(round(c, 2) for c in (span.get("color") or ())[:3]) == (1.0, 1.0, 1.0) and not has_images:
                kinds["white"].append(text)
            elif span.get("size", 12) * abs(span.get("scale", 1) or 1) < 1.0:
                kinds["tiny"].append(text)
            elif not bbox.is_empty and not bbox.intersects(prect):
                kinds["offpage"].append(text)
    return kinds


def privacy_scan(pdf: bytes) -> Dict[str, Any]:
    """Looks for personal or hidden information. Returns {"score", "findings", "found_terms"}.
    A finding's "fix" says which tool cleans it: "remove-metadata" (strip_private_data),
    "redact" (redact found_terms), "flatten", or None (needs a manual look)."""
    doc = _open_pdf(pdf)
    findings: List[Dict[str, Any]] = []

    def add(kind, severity, title, detail, fix):
        findings.append({"kind": kind, "severity": severity, "title": title, "detail": detail, "fix": fix})

    try:
        meta = doc.metadata or {}
        author = (meta.get("author") or "").strip()
        if author:
            add("metadata-author", "medium", "Author name in document properties",
                f"The file says it was written by \u201c{author}\u201d.", "remove-metadata")
        other = []
        for key, label in (("title", "title"), ("subject", "subject"), ("keywords", "keywords"),
                           ("creator", "created with"), ("producer", "produced by")):
            v = (meta.get(key) or "").strip()
            if v:
                other.append(f"{label}: {v}")
        for key, label in (("creationDate", "created"), ("modDate", "modified")):
            v = (meta.get(key) or "").strip()
            if v:
                m = re.match(r"D:(\d{4})(\d{2})(\d{2})", v)
                other.append(f"{label}: {m.group(1)}-{m.group(2)}-{m.group(3)}" if m else f"{label}: {v}")
        if other:
            add("metadata", "low", "Document properties",
                "The file carries details about its origin: " + "; ".join(other)[:500] + ".", "remove-metadata")
        try:
            xmp = doc.get_xml_metadata() or ""
        except Exception:
            xmp = ""
        if xmp.strip():
            history = "xmpMM:History" in xmp or "stEvt:" in xmp
            add("xmp", "medium" if history else "low", "Hidden XMP metadata",
                "An XML metadata packet is embedded" + (" including an editing history" if history else "")
                + ". It can hold names, software and earlier titles.", "remove-metadata")

        actions = _actions_in(pdf)
        if actions["actions"] or actions["open_action"]:
            parts = [f"{n} {label} action{'s' if n > 1 else ''}" for label, n in actions["actions"].items()]
            if actions["open_action"]:
                parts.append("an action that runs automatically when the file is opened")
            add("javascript", "high", "Scripts or automatic actions",
                "Found " + ", ".join(parts) + ". These can track when the file is opened or run code.", "remove-metadata")

        names = []
        try:
            for i in range(doc.embfile_count()):
                names.append(doc.embfile_info(i).get("filename") or f"file {i + 1}")
        except Exception:
            pass
        annot_count = 0
        annot_authors = set()
        for page in doc:
            for a in page.annots() or []:
                if a.type[0] == pymupdf.PDF_ANNOT_FILE_ATTACHMENT:
                    try:
                        names.append(a.file_info.get("filename") or "attachment")
                    except Exception:
                        names.append("attachment")
                    continue
                if a.type[0] in (pymupdf.PDF_ANNOT_LINK, pymupdf.PDF_ANNOT_WIDGET, pymupdf.PDF_ANNOT_POPUP):
                    continue
                annot_count += 1
                who = (a.info.get("title") or "").strip()
                if who:
                    annot_authors.add(who)
        if names:
            add("attachments", "high", "Attached files",
                f"{len(names)} file(s) are hidden inside this PDF: " + ", ".join(names[:10]) + ".", "remove-metadata")
        if annot_count:
            detail = f"{annot_count} comment(s) or markup annotation(s)"
            if annot_authors:
                detail += " written by " + ", ".join(sorted(annot_authors)[:10])
            add("annotations", "medium" if annot_authors else "low", "Comments and annotations", detail + ".", "flatten")

        filled = []
        for page in doc:
            for w in page.widgets() or []:
                v = w.field_value
                if v not in (None, "", False, "Off") and w.field_type not in (pymupdf.PDF_WIDGET_TYPE_BUTTON,):
                    filled.append(w.field_name or "field")
        if filled:
            add("form-fields", "medium", "Filled-in form fields",
                f"{len(filled)} form field(s) contain values (e.g. {', '.join(filled[:5])}). Anyone can read or change them.",
                "flatten")

        hidden = _hidden_text(doc)
        labels = {"invisible": "invisible text", "white": "white text", "tiny": "text too small to see",
                  "offpage": "text outside the visible page"}
        hidden_parts = [f"{labels[k]} (\u201c{' '.join(v)[:60]}\u201d)" for k, v in hidden.items() if v]
        if hidden_parts:
            add("hidden-text", "medium", "Hidden text", "Found " + "; ".join(hidden_parts) + ".", None)

        text = "\n".join(page.get_text() for page in doc)
        sensitive = _find_sensitive(text)
        labels2 = {"email": ("Email addresses", "medium"), "phone": ("Phone numbers", "medium"),
                   "card": ("Credit card numbers", "high"), "iban": ("Bank account numbers (IBAN)", "high")}
        for key, (title, sev) in labels2.items():
            vals = sensitive[key]
            if vals:
                shown = vals[:5] if key in ("email", "phone") else [v[:4] + "\u2026" + v[-4:] for v in vals[:5]]
                add(key, sev, title, f"{len(vals)} found: " + ", ".join(shown) + (" \u2026" if len(vals) > 5 else ""), "redact")

        if actions["incremental"] > 0:
            add("history", "medium", "Earlier versions are recoverable",
                f"This file was edited {actions['incremental']} time(s) by appending changes. "
                "The previous versions, including anything deleted, can still be recovered.", "remove-metadata")

        uris = set(actions["uris"])
        for page in doc:
            for link in page.get_links():
                if link.get("kind") == pymupdf.LINK_URI and link.get("uri"):
                    uris.add(link["uri"])
        external = sorted(u for u in uris if re.match(r"(https?|ftp)://", u, re.IGNORECASE))
        if external:
            domains = sorted({re.sub(r"^[a-z]+://([^/:?#]+).*$", r"\1", u, flags=re.IGNORECASE) for u in external})
            add("links", "low", "Links to websites",
                f"{len(external)} link(s) to: " + ", ".join(domains[:10]) + ". Clicking them can reveal that the file was opened.",
                None)
    finally:
        doc.close()

    score = max(0, 100 - sum(SEVERITY_PENALTY[f["severity"]] for f in findings))
    order = {"high": 0, "medium": 1, "low": 2}
    findings.sort(key=lambda f: order[f["severity"]])
    terms: List[str] = []
    for key in ("email", "phone", "card", "iban"):
        terms += sensitive.get(key, [])
    return {"score": score, "findings": findings, "found_terms": list(dict.fromkeys(terms))[:50]}


def strip_private_data(pdf: bytes) -> bytes:
    """Removes metadata, XMP, scripts and automatic actions, attachments and annotation authors,
    and rewrites the file in full so earlier versions can't be recovered."""
    _open_pdf(pdf).close()
    import pikepdf
    try:
        p = pikepdf.open(io.BytesIO(pdf))
    except Exception:
        raise ValueError("This PDF could not be read. Try the Repair tool first.")
    risky = set(_ACTION_TYPES) | {"/Rendition", "/Sound", "/Movie", "/RichMediaExecute"}
    with p:
        root = p.Root
        for key in ("/Metadata", "/OpenAction", "/AA", "/AF", "/Collection", "/PieceInfo", "/SpiderInfo"):
            if key in root:
                del root[key]
        names = root.get("/Names")
        if isinstance(names, pikepdf.Dictionary):
            for key in ("/JavaScript", "/EmbeddedFiles"):
                if key in names:
                    del names[key]
        acro = root.get("/AcroForm")
        if isinstance(acro, pikepdf.Dictionary) and "/XFA" in acro:
            del acro["/XFA"]
        if "/Info" in p.trailer:
            del p.trailer["/Info"]

        def clean_action_dict(d):
            if "/AA" in d:
                del d["/AA"]
            a = d.get("/A")
            if isinstance(a, pikepdf.Dictionary) and str(a.get("/S", "")) in risky:
                del d["/A"]

        for page in p.pages:
            pobj = page.obj
            for key in ("/AA", "/Metadata", "/PieceInfo"):
                if key in pobj:
                    del pobj[key]
            annots = pobj.get("/Annots")
            if isinstance(annots, pikepdf.Array):
                keep = pikepdf.Array()
                for annot in annots:
                    if isinstance(annot, pikepdf.Dictionary) and annot.get("/Subtype") == "/FileAttachment":
                        continue
                    keep.append(annot)
                pobj.Annots = keep
        for obj in p.objects:
            if not isinstance(obj, pikepdf.Dictionary):
                continue
            clean_action_dict(obj)
            if str(obj.get("/S", "")) in risky:
                for key in ("/JS", "/F", "/Win", "/Mac", "/Unix", "/Next"):
                    if key in obj:
                        del obj[key]
            if "/Subtype" in obj and obj.get("/Subtype") != "/Widget" and "/Rect" in obj:
                # Markup annotation: drop the author name and modification time.
                for key in ("/T", "/M", "/CreationDate", "/NM"):
                    if key in obj:
                        del obj[key]
            for key in ("/PieceInfo", "/LastModified"):
                if key in obj:
                    del obj[key]
            if "/Metadata" in obj:
                del obj["/Metadata"]
        out = io.BytesIO()
        # A full save writes only reachable objects: old revisions and orphaned data are gone.
        p.save(out, compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate,
               fix_metadata_version=False, deterministic_id=True)
    return out.getvalue()


# ---------------------------------------------------------------------------
# 8. PDF/A
# ---------------------------------------------------------------------------

def find_ghostscript() -> Optional[str]:
    for name in ("gswin64c", "gswin32c", "gs"):
        exe = shutil.which(name)
        if exe:
            return exe
    if os.name == "nt":
        for root in _windows_program_dirs():
            for exe in ("gswin64c.exe", "gswin32c.exe"):
                hits = sorted(glob.glob(os.path.join(root, "gs", "gs*", "bin", exe)), reverse=True)
                if hits:
                    return hits[0]
    return None


def _find_srgb_profile(gs: str) -> Optional[str]:
    candidates = []
    for mod in ("pikepdf", "ocrmypdf"):
        try:
            m = __import__(mod)
            base = os.path.dirname(m.__file__)
            candidates += [os.path.join(base, "pdfa", "data", "sRGB.icc"), os.path.join(base, "data", "sRGB.icc")]
        except Exception:
            pass
    gs_root = os.path.dirname(os.path.dirname(os.path.realpath(gs)))
    candidates += [os.path.join(gs_root, "iccprofiles", "srgb.icc"),
                   "/usr/share/color/icc/ghostscript/srgb.icc", "/usr/share/ghostscript/iccprofiles/srgb.icc",
                   "/usr/share/color/icc/sRGB.icc"]
    candidates += sorted(glob.glob("/usr/share/ghostscript/*/iccprofiles/srgb.icc"), reverse=True)
    for c in candidates:
        if os.path.isfile(c):
            return c
    return None


def _ps_string(path: str) -> str:
    p = Path(path).resolve().as_posix()
    return "(" + p.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ")"


def pdf_to_pdfa(pdf: bytes) -> bytes:
    """Converts to PDF/A-2b (archival) with Ghostscript."""
    gs = find_ghostscript()
    if not gs:
        raise ToolUnavailableError("PDF/A conversion needs Ghostscript on the server.")
    _open_pdf(pdf).close()
    workdir = tempfile.mkdtemp(prefix="ultrapdf_pdfa_")
    try:
        src = os.path.join(workdir, "in.pdf")
        dst = os.path.join(workdir, "out.pdf")
        with open(src, "wb") as f:
            f.write(pdf)
        cmd = [gs, "-dSAFER", "-dPDFA=2", "-dBATCH", "-dNOPAUSE", "-dQUIET", "-dNOOUTERSAVE", "-dPDFACompatibilityPolicy=1",
               "-sDEVICE=pdfwrite", "-sColorConversionStrategy=RGB", "-sProcessColorModel=DeviceRGB",
               "-dEmbedAllFonts=true", "-dSubsetFonts=true", "-dAutoRotatePages=/None",
               f"-sOutputFile={dst}"]
        icc = _find_srgb_profile(gs)
        if icc:
            icc_copy = os.path.join(workdir, "srgb.icc")
            shutil.copyfile(icc, icc_copy)
            pdfa_def = os.path.join(workdir, "PDFA_def.ps")
            with open(pdfa_def, "w", encoding="ascii") as f:
                f.write("%!\n"
                        f"/ICCProfile {_ps_string(icc_copy)} def\n"
                        "[/_objdef {icc_PDFA} /type /stream /OBJ pdfmark\n"
                        "[{icc_PDFA} << /N 3 >> /PUT pdfmark\n"
                        "[{icc_PDFA} ICCProfile (r) file /PUT pdfmark\n"
                        "[/_objdef {OutputIntent_PDFA} /type /dict /OBJ pdfmark\n"
                        "[{OutputIntent_PDFA} << /Type /OutputIntent /S /GTS_PDFA1 /DestOutputProfile {icc_PDFA}"
                        " /OutputConditionIdentifier (sRGB) /Info (sRGB IEC61966-2.1) >> /PUT pdfmark\n"
                        "[{Catalog} << /OutputIntents [ {OutputIntent_PDFA} ] >> /PUT pdfmark\n")
            cmd += [f"--permit-file-read={Path(workdir).resolve().as_posix()}/", "-sOutputICCProfile=" + icc_copy, pdfa_def]
        cmd.append(src)
        result = _run(cmd, timeout=180, cwd=workdir)
        if result.returncode != 0 or not os.path.isfile(dst) or os.path.getsize(dst) == 0:
            detail = (result.stderr or result.stdout or b"").decode("utf-8", "replace").strip().splitlines()
            logging.getLogger(__name__).warning("Ghostscript PDF/A failed: %s", detail[-3:])
            raise ValueError("This PDF could not be converted to PDF/A. It may be damaged; try the Repair tool first.")
        with open(dst, "rb") as f:
            out = f.read()
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    if _page_count(out) <= 0:
        raise ValueError("This PDF could not be converted to PDF/A.")
    return out
